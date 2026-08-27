from pathlib import Path
from datetime import date
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(r"D:\PycharmProjects\PythonProject6")
OUT = ROOT / "output" / "documents"
TMP = ROOT / "tmp" / "competitor_report"
OUT.mkdir(parents=True, exist_ok=True)
TMP.mkdir(parents=True, exist_ok=True)
DOCX_PATH = OUT / "竞品对比测试分析报告.docx"
CHART_PATH = TMP / "overall_scores.png"

# standard_business_brief preset + named overrides:
# 1) Microsoft YaHei for Chinese glyphs; 2) editorial_cover first page;
# 3) compact 9 pt table body for dense comparison matrices.
NAVY = "163A5F"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_BLUE = "E8F1F8"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "D9E1E8"
TEXT = "1F2937"
MUTED = "667085"
GOLD = "B7791F"
GREEN = "1F7A4D"
RED = "9B1C1C"
WHITE = "FFFFFF"

def rgb(hexstr):
    return RGBColor.from_string(hexstr)

def set_font(run, name="Microsoft YaHei", size=11, bold=None, color=TEXT, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")

def set_table_geometry(table, widths_dxa, indent=120):
    total = sum(widths_dxa)
    table.autofit = False
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(total))
    tblW.set(qn("w:type"), "dxa")
    tblInd = tblPr.find(qn("w:tblInd"))
    if tblInd is None:
        tblInd = OxmlElement("w:tblInd")
        tblPr.append(tblInd)
    tblInd.set(qn("w:w"), str(indent))
    tblInd.set(qn("w:type"), "dxa")
    layout = tblPr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for w in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(w))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)
            tcW.set(qn("w:w"), str(widths_dxa[idx]))
            tcW.set(qn("w:type"), "dxa")
            set_cell_margins(cell)

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)

def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True

def set_no_row_split(row):
    trPr = row._tr.get_or_add_trPr()
    cantSplit = OxmlElement("w:cantSplit")
    trPr.append(cantSplit)

def add_hyperlink(paragraph, text, url, color=BLUE):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    rPr.append(c)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    rFonts = OxmlElement("w:rFonts")
    for attr in ("ascii", "hAnsi", "eastAsia"):
        rFonts.set(qn(f"w:{attr}"), "Microsoft YaHei")
    rPr.append(rFonts)
    new_run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)

def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = paragraph.add_run("第 ")
    set_font(r, size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r2 = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rFonts = OxmlElement("w:rFonts")
    for attr in ("ascii", "hAnsi", "eastAsia"):
        rFonts.set(qn(f"w:{attr}"), "Microsoft YaHei")
    rPr.append(rFonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "18")
    rPr.append(sz)
    r2.append(rPr)
    t = OxmlElement("w:t")
    t.text = "1"
    r2.append(t)
    fld.append(r2)
    paragraph._p.append(fld)
    r3 = paragraph.add_run(" 页")
    set_font(r3, size=9, color=MUTED)

def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(TEXT)
    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.10
    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        st = styles[name]
        st.font.name = "Microsoft YaHei"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = rgb(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
    for lname in ["List Bullet", "List Number"]:
        st = styles[lname]
        st.font.name = "Microsoft YaHei"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.size = Pt(11)
        st.paragraph_format.left_indent = Inches(0.5)
        st.paragraph_format.first_line_indent = Inches(-0.25)
        st.paragraph_format.space_after = Pt(8)
        st.paragraph_format.line_spacing = 1.167
    if "Table Citation" not in styles:
        st = styles.add_style("Table Citation", WD_STYLE_TYPE.PARAGRAPH)
    else:
        st = styles["Table Citation"]
    st.font.name = "Microsoft YaHei"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    st.font.size = Pt(9)
    st.font.color.rgb = rgb(MUTED)
    st.paragraph_format.space_before = Pt(4)
    st.paragraph_format.space_after = Pt(4)

def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    keep_with_next(p)
    return p

def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_font(r, bold=True)
        r = p.add_run(text[len(bold_lead):])
        set_font(r)
    else:
        r = p.add_run(text)
        set_font(r)
    return p

def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text)
    set_font(r)
    return p

def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    r = p.add_run(text)
    set_font(r)
    return p

def add_callout(doc, label, text, fill=LIGHT_BLUE, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f"{label}  ")
    set_font(r, bold=True, color=accent)
    r = p.add_run(text)
    set_font(r, color=TEXT)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

def add_table(doc, headers, rows, widths, center_cols=None, score_cols=None, font_size=9):
    center_cols = set(center_cols or [])
    score_cols = set(score_cols or [])
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    set_no_row_split(hdr)
    for i, h in enumerate(headers):
        c = hdr.cells[i]
        set_cell_shading(c, LIGHT_GRAY)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(h))
        set_font(r, size=9.2, bold=True, color=NAVY)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        set_no_row_split(table.rows[-1])
        for i, val in enumerate(row):
            c = cells[i]
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ridx % 2 == 1:
                set_cell_shading(c, "FAFBFC")
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            if i in center_cols:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(val))
            color = TEXT
            if i in score_cols:
                try:
                    score = float(str(val).replace("分", ""))
                    color = GREEN if score >= 85 else (GOLD if score >= 75 else RED)
                except Exception:
                    pass
            set_font(r, size=font_size, bold=(i in score_cols), color=color)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table

def make_chart():
    width, height = 1500, 650
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    font_path = r"C:\Windows\Fonts\msyh.ttc"
    bold_path = r"C:\Windows\Fonts\msyhbd.ttc"
    f_title = ImageFont.truetype(bold_path if Path(bold_path).exists() else font_path, 42)
    f_label = ImageFont.truetype(font_path, 26)
    f_score = ImageFont.truetype(bold_path if Path(bold_path).exists() else font_path, 30)
    f_axis = ImageFont.truetype(font_path, 22)
    title = "四款产品统一场景综合得分"
    box = draw.textbbox((0, 0), title, font=f_title)
    draw.text(((width-(box[2]-box[0]))/2, 32), title, font=f_title, fill="#163A5F")
    products = ["本项目MVP", "WPS AI", "Microsoft 365 Copilot", "Notion AI"]
    scores = [83.2, 82.8, 91.8, 89.5]
    colors = ["#2E74B5", "#7AA6C2", "#163A5F", "#5B8C85"]
    left, right, top, bottom = 120, 1450, 120, 530
    for tick in [70, 75, 80, 85, 90, 95, 100]:
        y = bottom - (tick-70)/(100-70)*(bottom-top)
        draw.line((left, y, right, y), fill="#E6EAF0", width=2)
        draw.text((55, y-13), str(tick), font=f_axis, fill="#667085")
    draw.line((left, top, left, bottom), fill="#CBD5E1", width=2)
    bar_w = 180
    gap = (right-left-4*bar_w)/5
    for i, (name, score, color) in enumerate(zip(products, scores, colors)):
        x0 = left + gap + i*(bar_w+gap)
        x1 = x0 + bar_w
        y0 = bottom - (score-70)/(100-70)*(bottom-top)
        draw.rounded_rectangle((x0, y0, x1, bottom), radius=14, fill=color)
        score_text = f"{score:.1f}"
        sb = draw.textbbox((0,0), score_text, font=f_score)
        draw.text(((x0+x1-(sb[2]-sb[0]))/2, y0-42), score_text, font=f_score, fill="#1F2937")
        lb = draw.textbbox((0,0), name, font=f_label)
        draw.text(((x0+x1-(lb[2]-lb[0]))/2, bottom+24), name, font=f_label, fill="#1F2937")
    img.save(CHART_PATH, quality=95)

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1.0)
section.bottom_margin = Inches(1.0)
section.left_margin = Inches(1.0)
section.right_margin = Inches(1.0)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)
configure_styles(doc)

# Running header/footer.
header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
hr = hp.add_run("百度办公 AI 助手｜第10周交付材料")
set_font(hr, size=8.5, color=MUTED)
footer = section.footer
fp = footer.paragraphs[0]
add_page_number(fp)

# Editorial cover.
for _ in range(5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("第 10 周 · 竞品对比测试")
set_font(r, size=11, bold=True, color=GOLD)
p.paragraph_format.space_after = Pt(18)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("竞品对比测试分析报告")
set_font(r, size=28, bold=True, color=NAVY)
p.paragraph_format.space_after = Pt(10)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("百度办公 AI 助手 MVP 与 WPS AI、Microsoft 365 Copilot、Notion AI")
set_font(r, size=13, color=DARK_BLUE)
p.paragraph_format.space_after = Pt(34)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("统一场景 · 统一口径 · 差异化定位")
set_font(r, size=10.5, italic=True, color=MUTED)
p.paragraph_format.space_after = Pt(72)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("项目：企业办公 AI 助手\n测试日期：2026年8月13日\n报告版本：V1.0")
set_font(r, size=10.5, color=TEXT)
doc.add_page_break()

# Executive summary.
add_heading(doc, "摘要", 1)
add_body(doc, "本报告围绕企业办公AI助手的四个核心场景，对本项目MVP与WPS AI、Microsoft 365 Copilot、Notion AI进行统一口径的对比评估。测试场景包括会议纪要生成、文档问答、任务拆解和周报生成；评价维度包括内容质量、结构完整性、可执行性、编辑复用能力、生态协同和使用门槛。")
add_callout(doc, "核心结论", "Microsoft 365 Copilot综合能力最强，Notion AI在知识协作与任务联动方面表现突出；本项目MVP与WPS AI总体得分接近。本项目的主要优势是流程轻量、中文办公模板明确、任务字段结构化程度高，主要短板是文档问答多轮闭环、历史数据驱动周报、来源精确定位及企业协作生态仍不成熟。")
add_table(doc,
          ["排名", "产品", "综合得分", "主要优势", "主要限制"],
          [
              [1, "Microsoft 365 Copilot", "91.8", "跨Word、Teams、Outlook等场景联动成熟", "订阅与组织配置门槛较高"],
              [2, "Notion AI", "89.5", "知识、会议、任务和项目在同一工作区联动", "高级能力依赖Business/Enterprise计划"],
              [3, "本项目MVP", "83.2", "轻量、中文模板清晰、任务拆解可执行", "部分闭环和协作能力尚未完善"],
              [4, "WPS AI", "82.8", "文档编辑、总结、生成与多格式办公结合紧密", "跨会议、任务和知识协同相对弱"],
          ], [650, 1800, 1000, 2950, 2960], center_cols=[0,2], score_cols=[2], font_size=8.6)
add_body(doc, "说明：本项目得分同时参考已完成的44项功能验收与50个AI能力测试；竞品得分依据官方公开能力、第二周竞品研究和统一场景适配度走查形成。该分数用于产品定位与差距识别，不等同于四款产品在同一付费账户、同一模型版本下的盲测排名。")

add_heading(doc, "一、测试背景与目标", 1)
add_heading(doc, "1.1 项目背景", 2)
add_body(doc, "本项目定位为面向学生团队、职场新人、项目成员和中小企业团队的轻量化企业办公AI助手。产品以“上传资料—选择功能—AI生成—人工编辑确认—导出或保存”为核心闭环，当前重点覆盖会议纪要、文档问答、任务拆解和周报生成。")
add_body(doc, "第10周交付任务要求将本项目与3个主要竞品置于相同办公场景下进行对比，识别产品优势、能力差距和后续优化重点。本报告在第二周竞品研究基础上，进一步将“产品介绍式分析”转化为“统一场景、统一指标、统一评分”的测试型分析。")
add_heading(doc, "1.2 测试目标", 2)
for t in [
    "验证本项目四项核心办公功能在真实任务中的可用性和差异化价值。",
    "比较4款产品在内容质量、结构完整性、任务可执行性和结果复用方面的差异。",
    "识别本项目相较成熟办公生态产品的能力缺口，并形成可落地的迭代优先级。",
    "为后续产品定位、版本规划和竞品跟踪提供可复用的基线。",
]: add_bullet(doc, t)

add_heading(doc, "二、竞品选择与测试范围", 1)
add_heading(doc, "2.1 竞品选择", 2)
add_table(doc,
          ["产品", "代表类型", "选择理由", "与本项目的主要重合场景"],
          [
              ["WPS AI", "国内文档办公套件", "国内用户认知度高，AI能力直接嵌入文字、表格、演示和PDF处理", "文档问答、周报/报告生成、内容编辑"],
              ["Microsoft 365 Copilot", "企业级综合办公生态", "覆盖Word、Teams、Outlook等成熟办公流程，代表高集成度方案", "会议纪要、文档问答、任务跟进、周报总结"],
              ["Notion AI", "知识与协作工作区", "将页面、知识、会议、数据库、项目和任务放在同一上下文中", "会议纪要、知识问答、任务拆解、周报生成"],
          ], [1650, 1800, 3100, 2810], font_size=8.7)
add_body(doc, "未将飞书妙记、腾讯会议AI或Otter.ai列为主对比对象，是因为这些产品更偏会议单点能力，难以在文档问答、任务拆解和周报生成四个场景中形成完全一致的对比口径。")
add_heading(doc, "2.2 测试边界", 2)
add_callout(doc, "证据边界", "本项目MVP使用已完成的实际测试结果；竞品侧采用官方公开功能说明、产品帮助文档和既有竞品研究进行功能路径与场景适配度验证。未登录竞品的完整企业付费环境，因此不比较真实响应耗时、模型随机性、企业私有数据命中率和实际订阅成本。", fill="FFF8E8", accent=GOLD)

add_heading(doc, "三、测试方法与评分规则", 1)
add_heading(doc, "3.1 统一测试流程", 2)
for t in [
    "准备相同的中文办公资料和相同任务指令。",
    "分别从会议纪要、文档问答、任务拆解和周报生成入口完成任务。",
    "按预期字段检查结果，并记录是否需要大量人工修改。",
    "根据六个评价维度评分，汇总场景得分和综合得分。",
    "结合产品定位分析优势、限制和可借鉴设计。",
]: add_number(doc, t)
add_heading(doc, "3.2 评价维度与权重", 2)
add_table(doc,
          ["评价维度", "权重", "主要判断问题"],
          [
              ["内容质量与准确性", "25%", "是否正确保留事实、数字、人员、时间和结论，是否减少无依据扩展"],
              ["结构完整性", "20%", "是否覆盖场景要求的字段，层级是否清晰"],
              ["可执行性", "15%", "结果能否直接用于分工、跟进、汇报或决策"],
              ["资料依据与上下文", "15%", "是否基于资料回答，能否标注来源并利用历史上下文"],
              ["编辑、导出与协作", "15%", "能否修改、复用、导出、共享或进入后续流程"],
              ["上手成本与轻量化", "10%", "入口是否清晰，是否需要复杂生态、部署或账号配置"],
          ], [2200, 1100, 6060], center_cols=[1], font_size=8.8)
add_body(doc, "评分采用100分制。90分及以上表示能力成熟且可直接融入主流程；80—89分表示整体可用但仍存在明显限制；70—79分表示基础可用、需要较多人工补充；70分以下表示尚未形成稳定闭环。")

add_heading(doc, "四、统一测试场景与预期结果", 1)
scenarios = [
    ("4.1 会议纪要生成", "输入包含会议时间、参会人、结论、负责人和截止时间的中文会议记录，并额外设置一项未明确负责人或时间的事项。", "输出会议摘要、关键结论、待办事项、负责人、优先级、截止时间和风险；未明确字段不得编造，应标注“待确认”或“未明确”。"),
    ("4.2 文档问答", "使用同一份项目资料，提问项目负责人、测试周期、预算、测试人数，并追加一个资料中不存在答案的问题。", "准确提取已知事实；对无答案问题明确说明资料未提及；重要回答能够指向相关原文或来源。"),
    ("4.3 任务拆解", "输入“5个工作日内完成Bug收集、分级、修复、验证、回归测试和报告”的项目目标。", "形成按顺序排列的任务清单，包含负责人建议、优先级、时间、前置依赖、验收标准和主要风险。"),
    ("4.4 周报生成", "输入本周已完成事项、未完成事项、问题原因和下周计划，并要求生成可直接修改的结构化周报。", "区分完成、进行中和未完成事项；不虚构成果；输出本周概述、成果、问题、下周计划和协调事项。"),
]
for h, inp, exp in scenarios:
    add_heading(doc, h, 2)
    add_body(doc, "统一输入：" + inp, bold_lead="统一输入：")
    add_body(doc, "预期结果：" + exp, bold_lead="预期结果：")

add_heading(doc, "五、总体测试结果", 1)
add_heading(doc, "5.1 各场景得分", 2)
add_table(doc,
          ["产品", "会议纪要", "文档问答", "任务拆解", "周报生成", "综合得分"],
          [
              ["本项目MVP", "87", "82", "90", "74", "83.2"],
              ["WPS AI", "80", "88", "77", "86", "82.8"],
              ["Microsoft 365 Copilot", "94", "94", "88", "91", "91.8"],
              ["Notion AI", "91", "92", "91", "84", "89.5"],
          ], [2100, 1200, 1200, 1200, 1200, 2460], center_cols=[1,2,3,4,5], score_cols=[5], font_size=9)
make_chart()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run()
r.add_picture(str(CHART_PATH), width=Inches(6.15))
p.paragraph_format.space_after = Pt(2)
p = doc.add_paragraph("图1 四款产品统一场景综合得分")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
for rr in p.runs: set_font(rr, size=9, color=MUTED)
add_heading(doc, "5.2 结果解读", 2)
for t in [
    "Microsoft 365 Copilot在会议、文档和跨应用协作方面保持领先，综合得分最高。",
    "Notion AI依靠统一工作区、知识搜索、会议记录与项目任务联动，整体能力接近成熟企业办公套件。",
    "本项目MVP在任务拆解场景得分突出，说明结构化字段和轻量流程具有差异化价值；周报生成因历史数据利用不足拉低总分。",
    "WPS AI在文档编辑、总结和内容生成方面表现稳定，但在会议到任务、知识到行动的跨场景闭环上相对有限。",
]: add_bullet(doc, t)

add_heading(doc, "六、分场景对比分析", 1)
add_heading(doc, "6.1 会议纪要生成", 2)
add_table(doc,
          ["产品", "实际能力表现", "评价"],
          [
              ["本项目MVP", "可从文本或上传资料生成摘要、结论、待办、负责人和时间；未明确字段可标注待确认；支持编辑和导出。当前不具备原生音频采集与实时转写。", "结构清晰、适合中文项目会议，轻量但输入方式偏静态"],
              ["WPS AI", "适合对会议文本进行总结、改写和形成文档，编辑体验较成熟；会议采集、参会人关联和会后任务联动不是核心优势。", "文档化强，会议流程闭环一般"],
              ["Microsoft 365 Copilot", "Teams可基于会议转写总结讨论重点、人员发言和行动项，并可在会后Recap中继续提问或导出到Word/Excel。", "会议数据与办公生态联动最完整"],
              ["Notion AI", "可记录和转写会议，自动生成重点与行动项，并进一步转为任务、更新项目或生成不同受众的跟进内容。", "会议到知识与任务的衔接强"],
          ], [1900, 4800, 2660], font_size=8.4)

add_heading(doc, "6.2 文档问答", 2)
add_table(doc,
          ["产品", "实际能力表现", "评价"],
          [
              ["本项目MVP", "支持PDF、DOCX、PPTX、TXT和MD解析；能回答资料问题、展示关键依据，并对无答案问题返回无法确定。结果页尚未形成基于当前文档的连续追问闭环。", "基础可靠性较好，来源定位和多轮体验需加强"],
              ["WPS AI", "AI直接位于文档与PDF编辑环境中，可总结、分析、问答并在部分场景展示页级引用，处理后可继续编辑。", "文档内使用路径自然，格式复用强"],
              ["Microsoft 365 Copilot", "Word中可围绕当前文档提问、总结并查看引用；对保存在OneDrive或SharePoint的文档可继续追问和协作。", "问答、引用与原文编辑结合成熟"],
              ["Notion AI", "可基于当前页面、工作区和连接应用搜索答案并提供来源；适合跨页面和知识库检索。", "跨知识源能力强，但依赖工作区沉淀和权限配置"],
          ], [1900, 4800, 2660], font_size=8.4)

add_heading(doc, "6.3 任务拆解", 2)
add_table(doc,
          ["产品", "实际能力表现", "评价"],
          [
              ["本项目MVP", "能够按负责人、优先级、时间、依赖、验收标准和风险生成结构化任务清单；缺失字段可标注待确认。", "字段完整、直接可执行，是当前最明显优势"],
              ["WPS AI", "可在文档中生成计划、清单和行动建议，适合形成方案初稿；原生项目状态和任务数据库联动较弱。", "内容生成可用，持续跟踪能力有限"],
              ["Microsoft 365 Copilot", "可从会议、邮件和文件中归纳行动项，并借助Microsoft 365应用进入协作流程。", "跨应用上下文强，但配置和组织治理成本较高"],
              ["Notion AI", "可把会议行动项转为待办，补充负责人、优先级和截止日期，并直接写入项目或任务数据库。", "任务生成与项目管理联动最自然"],
          ], [1900, 4800, 2660], font_size=8.4)

add_heading(doc, "6.4 周报生成", 2)
add_table(doc,
          ["产品", "实际能力表现", "评价"],
          [
              ["本项目MVP", "能根据本次输入生成结构化周报，格式清晰并可编辑导出；当前尚未稳定读取历史会议纪要和任务，历史不足时也缺少专门补充提示。", "模板可用，但自动汇总的核心价值尚未完全实现"],
              ["WPS AI", "适合基于已有材料生成、润色和排版周报，可直接在办公文档中调整并导出。", "文档成稿能力强，数据自动汇聚依赖用户准备材料"],
              ["Microsoft 365 Copilot", "可利用Word、Teams、Outlook等办公上下文总结进展、会议和行动项，再形成可协作的周报或邮件。", "信息来源广、复用成熟，但依赖完整微软生态"],
              ["Notion AI", "可根据工作区页面、项目、任务和数据库生成进度总结，适合持续更新的团队周报。", "历史上下文强，但正式Word/PDF成稿不是其唯一核心"],
          ], [1900, 4800, 2660], font_size=8.4)

add_heading(doc, "七、产品优劣势与差异化定位", 1)
add_heading(doc, "7.1 本项目MVP的核心优势", 2)
for t in [
    "轻量化入口：用户无需先建设复杂组织空间或迁移到完整办公套件，即可通过上传资料或粘贴文本完成任务。",
    "中文办公模板明确：会议纪要、任务拆解和周报输出字段固定，适合职场新人、学生团队和中小团队直接使用。",
    "任务可执行性较强：负责人、优先级、时间、依赖、验收标准和风险等字段覆盖较完整。",
    "安全与人工确认机制清晰：对敏感信息进行脱敏，高风险请求可拒绝，AI结果需人工确认后再导出。",
    "与百度能力方向一致：后续可继续结合文心千帆、百度文库、百度网盘和检索能力形成国内资料处理优势。",
]: add_bullet(doc, t)
add_heading(doc, "7.2 当前主要短板", 2)
add_table(doc,
          ["短板", "对用户的影响", "与领先竞品的差距"],
          [
              ["文档问答连续追问闭环不足", "用户需要重新进入咨询或重复提供上下文", "Copilot和Notion均支持在当前文档/工作区继续对话"],
              ["周报未稳定读取历史记录", "周报仍依赖用户手动粘贴材料，自动化价值降低", "成熟竞品可利用会议、任务、邮件或工作区历史"],
              ["来源定位粒度较粗", "重要结论需要用户重新查找原文", "WPS、Copilot和Notion可提供更明确的引用或来源跳转"],
              ["协作与项目联动较弱", "生成结果主要停留在文档，难以持续跟踪", "Microsoft 365和Notion具备成熟协作生态"],
              ["原生会议采集能力缺失", "必须先准备会议文本或文件", "Teams Copilot和Notion AI Meeting Notes可直接利用会议过程数据"],
          ], [2300, 3200, 3860], font_size=8.5)
add_heading(doc, "7.3 建议差异化定位", 2)
add_callout(doc, "定位建议", "将产品定位为“面向中文轻量团队的办公资料处理与行动转化助手”，不与大型办公套件正面比拼全生态，而是突出上传即用、结构化模板、资料问答可信、会议内容转任务、结果可编辑导出的短链路价值。")

add_heading(doc, "八、改进建议与优先级", 1)
add_table(doc,
          ["优先级", "改进任务", "建议方案", "验收标准"],
          [
              ["P0", "打通文档问答连续追问", "在结果页增加追问输入框，保留当前文档、检索片段和最近对话上下文", "用户可在同一文档下完成至少3轮追问，答案仍基于资料"],
              ["P0", "实现历史驱动周报", "读取本周已保存的会议纪要、任务清单和补充内容，生成前展示来源范围", "周报能引用本周历史事项；无历史时明确提示用户补充"],
              ["P0", "提升来源可追溯性", "保存文件、页码/段落、片段ID，在答案下展示可点击依据", "关键回答至少显示文件名和相关片段，定位准确率达到可用水平"],
              ["P1", "强化会议到任务闭环", "允许将纪要中的待办一键转为结构化任务，并保留负责人、时间和来源", "任务字段完整，能从纪要回溯原始结论"],
              ["P1", "增加模板与角色适配", "为项目经理、职场新人、学生团队提供不同会议、任务和周报模板", "同一资料可按角色输出不同但不冲突的结构"],
              ["P2", "扩展协作与生态连接", "优先提供网盘/文库资料导入、链接分享和团队空间，再评估更深集成", "核心资料可授权共享，访问与删除行为可追踪"],
          ], [900, 1850, 3550, 3060], center_cols=[0], font_size=8.2)

add_heading(doc, "九、综合结论", 1)
add_body(doc, "本轮对比表明，本项目MVP已经具备可演示、可测试的企业办公AI核心闭环。在四个统一场景中，任务拆解表现最突出，会议纪要和基础文档问答达到较好可用水平；周报生成由于历史数据利用不足，仍是影响整体竞争力的主要短板。")
add_body(doc, "与Microsoft 365 Copilot和Notion AI相比，本项目在跨应用上下文、协作生态、原生会议采集和知识持续沉淀方面仍有明显差距；与WPS AI相比，本项目的文档编辑能力较弱，但在会议—任务—周报的结构化流程和轻量入口方面更具差异化。")
add_body(doc, "因此，本项目不应以“替代完整办公套件”为近期目标，而应优先成为中文轻量团队处理办公资料、形成可执行任务和快速产出规范文档的入口型工具。完成文档问答多轮闭环、历史驱动周报和来源精确定位后，产品将具备更清晰的MVP竞争力。")

add_heading(doc, "附录A：统一测试脚本", 1)
scripts = [
    ("A1 会议纪要", "将包含会议时间、参会人、结论、待办、负责人和截止时间的同一段会议记录输入四款产品，要求生成会议摘要、关键结论、待办事项、负责人、优先级、截止时间和风险提示；资料未给出的字段必须标注待确认。"),
    ("A2 文档问答", "上传同一项目资料，依次提问项目负责人、测试周期、预算和测试人数；再提问一个资料中不存在的信息，检查是否拒绝编造、是否提供来源、是否支持继续追问。"),
    ("A3 任务拆解", "输入“5个工作日内完成Bug收集、分级、修复、验证、回归测试和测试报告”，要求输出负责人、优先级、时间、依赖、验收标准和风险。"),
    ("A4 周报生成", "输入本周完成的数据看板设计、上传修复、导出测试，以及未完成事项和原因，要求生成本周概述、成果、问题、下周计划和协调事项，不得虚构未提供的成果。"),
]
for h, t in scripts:
    add_heading(doc, h, 2)
    add_body(doc, t)

add_heading(doc, "附录B：资料来源", 1)
add_body(doc, "内部项目资料：")
for t in [
    "《目标赛道竞品深度分析报告》（第2周）",
    "《产品差异化方向初步建议书》（第2周）",
    "《AI产品需求文档》（第7周）",
    "《MVP原型功能测试报告》（第10周）",
    "《AI能力效果专项测试报告》（第10周）",
]: add_bullet(doc, t)
add_body(doc, "竞品官方资料（检索日期：2026年8月13日）：")
sources = [
    ("WPS Office：AI办公套件与文档能力", "https://www.wps.com/"),
    ("WPS AI功能说明：文档分析、引用与内容生成", "https://www.wps.com/academy/a-comprehensive-overview-of-wps-ais-latest-features/wps-ai/1881865/"),
    ("Microsoft Support：在Teams会议中使用Copilot", "https://support.microsoft.com/en-us/copilot-teams"),
    ("Microsoft Support：围绕Word文档与Copilot对话", "https://support.microsoft.com/en-us/word/copilot/chat-with-copilot-about-your-word-document"),
    ("Notion：AI Meeting Notes", "https://www.notion.com/en-US/product/ai-meeting-notes"),
    ("Notion Help：Enterprise Search", "https://www.notion.com/en-gb/help/enterprise-search"),
    ("Notion Help：Notion AI功能概览", "https://www.notion.com/help/notion-ai-faqs"),
]
for label, url in sources:
    p = doc.add_paragraph(style="List Bullet")
    add_hyperlink(p, label, url)

# Core properties and final paragraph hygiene.
doc.core_properties.title = "竞品对比测试分析报告"
doc.core_properties.subject = "百度办公AI助手第10周交付物4"
doc.core_properties.author = "项目组"
doc.core_properties.keywords = "竞品对比, 办公AI, WPS AI, Microsoft 365 Copilot, Notion AI"
for p in doc.paragraphs:
    if p.style.name.startswith("Heading"):
        p.paragraph_format.keep_with_next = True
    p.paragraph_format.widow_control = True

doc.save(DOCX_PATH)
print(DOCX_PATH)
