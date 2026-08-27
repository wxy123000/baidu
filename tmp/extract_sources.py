from pathlib import Path
from docx import Document
import glob
import pdfplumber

def show_docx(path, needles):
    print(f"\n=== {Path(path).name} ===")
    doc = Document(path)
    for p in doc.paragraphs:
        t = p.text.strip()
        if t and any(n in t for n in needles):
            print("P:", t[:900])
    for ti, table in enumerate(doc.tables):
        rows=[]
        for row in table.rows:
            vals=[c.text.strip().replace('\n',' / ') for c in row.cells]
            joined=' | '.join(vals)
            if any(n in joined for n in needles):
                rows.append(joined)
        if rows:
            print(f"TABLE {ti}")
            for x in rows[:30]: print(x[:1200])

week2 = glob.glob(r"C:\Users\20159\Desktop\实习\第二周\第二周\*.docx")
needles = ["WPS", "Microsoft", "Notion", "飞书", "差异化", "会议纪要", "文档问答", "任务", "周报", "核心结论", "竞品"]
for p in week2:
    if any(k in Path(p).name for k in ["竞品深度", "差异化方向", "标杆"]):
        show_docx(p, needles)

for p in glob.glob(r"C:\Users\20159\Desktop\实习\第十周\*.docx"):
    if Path(p).name.startswith("~$"):
        continue
    if any(k in Path(p).name for k in ["MVP原型功能测试报告", "AI能力效果专项测试报告"]):
        show_docx(p, ["通过率", "综合", "结论", "平均分", "50", "44", "36", "有条件", "准确性", "安全性", "有用性"])

prd = r"C:\Users\20159\Desktop\实习\第七周\AI 产品需求文档.pdf"
print(f"\n=== {Path(prd).name} ===")
with pdfplumber.open(prd) as pdf:
    for i, page in enumerate(pdf.pages, 1):
        text = page.extract_text() or ""
        if any(k in text for k in ["会议纪要", "文档问答", "任务拆解", "周报", "验收", "核心功能"]):
            print(f"PAGE {i}: {text[:1800].replace(chr(10),' / ')}")
