from docx import Document
from pathlib import Path

p=Path(r'D:\PycharmProjects\PythonProject6\tmp\competitor_review_copy.docx')
d=Document(p)
print('FILE',p,'SIZE',p.stat().st_size,'PARAS',len(d.paragraphs),'TABLES',len(d.tables))
print('\nHEADINGS')
for i,p in enumerate(d.paragraphs):
    if p.style.name.startswith('Heading') or (p.text.strip() and p.text.strip()[0:2] in ['一、','二、','三、','四、','五、','六、','七、','八、','九、']):
        print(i,p.style.name,repr(p.text.strip()))
print('\nKEY PARAS')
for i,p in enumerate(d.paragraphs):
    t=p.text.strip()
    if any(k in t for k in ['核心结论','说明：本项目得分','测试边界','评分采用','不等同','综合结论','本轮对比表明','未登录','得分']):
        print(i,repr(t))
print('\nTABLES')
for ti,t in enumerate(d.tables):
    print('\nTABLE',ti,'rows',len(t.rows),'cols',len(t.columns))
    for row in t.rows:
        print(' | '.join(c.text.strip().replace('\n',' / ') for c in row.cells))
