from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.oxml.ns import qn

path = Path(r"D:\PycharmProjects\PythonProject6\outputs\week11_simulated_dataset\模拟用户数据集构建说明.docx")
doc = Document(path)
headings = [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith("Heading")]
nonempty = [p.text for p in doc.paragraphs if p.text.strip()]
table_widths = []
for idx, table in enumerate(doc.tables, 1):
    grid = table._tbl.tblGrid
    widths = [int(col.get(qn("w:w"))) for col in grid]
    table_widths.append((idx, sum(widths), widths, len(table.rows), len(table.columns)))
with ZipFile(path) as z:
    names = z.namelist()
    media = [n for n in names if n.startswith("word/media/")]
    xml = z.read("word/document.xml").decode("utf-8", errors="replace")
    broken = z.testzip()
print({
    "file_bytes": path.stat().st_size,
    "paragraphs": len(doc.paragraphs),
    "nonempty_paragraphs": len(nonempty),
    "headings": headings,
    "tables": len(doc.tables),
    "table_widths": table_widths,
    "inline_shapes": len(doc.inline_shapes),
    "media": media,
    "sections": len(doc.sections),
    "page_size_inches": (round(doc.sections[0].page_width / 914400, 3), round(doc.sections[0].page_height / 914400, 3)),
    "margins_inches": tuple(round(x / 914400, 3) for x in (doc.sections[0].top_margin, doc.sections[0].right_margin, doc.sections[0].bottom_margin, doc.sections[0].left_margin)),
    "has_page_field": "PAGE" in xml,
    "zip_error": broken,
    "simulation_disclaimer_mentions": sum(t.count("模拟") for t in nonempty),
    "real_operation_mentions": sum(t.count("真实运营") for t in nonempty),
    "empty_table_cells": sum(1 for t in doc.tables for row in t.rows for cell in row.cells if not cell.text.strip()),
})
