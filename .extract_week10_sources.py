from docx import Document
from pathlib import Path
files=[
 r'C:\Users\20159\Desktop\实习\第十周\AI能力效果专项测试报告.docx',
 r'C:\Users\20159\Desktop\实习\第十周\MVP原型功能测试报告.docx',
 r'C:\Users\20159\Desktop\实习\第十周\竞品对比测试分析报告.docx',
 r'C:\Users\20159\Desktop\实习\第十周\数据复盘机制规范-完整版.docx',
 r'C:\Users\20159\Desktop\实习\第十周\项目核心数据看板原型设计.docx']
keys=['总结','结论','总体','通过率','通过','未通过','问题','建议','完成','交付','竞品','优势','不足','风险','测试结果','平均分','维度','原型','复盘']
for f in files:
 d=Document(f); print('\n\n#####',Path(f).name,'P',len(d.paragraphs),'T',len(d.tables))
 for i,p in enumerate(d.paragraphs):
  t=' '.join(p.text.split())
  if t and (p.style.name.startswith('Heading') or any(k in t for k in keys)):
   print(f'P{i:03d}[{p.style.name}] {t}')
 print('--- relevant table rows ---')
 for ti,t in enumerate(d.tables):
  for ri,row in enumerate(t.rows):
   s=' | '.join(' '.join(c.text.split()) for c in row.cells)
   if ri==0 or any(k in s for k in keys): print(f'T{ti+1}R{ri}: {s}')
