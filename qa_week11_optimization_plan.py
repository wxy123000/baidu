from pathlib import Path
from zipfile import ZipFile, BadZipFile
from docx import Document
from docx.oxml.ns import qn

p = Path(r"D:\PycharmProjects\PythonProject6\outputs\week11_simulated_dataset\项目AI产品全案优化迭代方案.docx")
issues = []
try:
    with ZipFile(p) as z:
        bad = z.testzip()
        if bad:
            issues.append(f"DOCX成员损坏: {bad}")
except BadZipFile:
    issues.append("文件不是有效DOCX")

doc = Document(p)
all_text = "\n".join(x.text for x in doc.paragraphs)
all_text += "\n" + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
required = [
    "产品功能", "AI融合", "用户体验", "产品增长", "V1.1 稳定性版",
    "总体结果返回率", "任务自动化", "安全与合规", "模拟数据",
    "90.1%", "81.2%", "68.4%", "P0", "验收与测试方案",
]
for term in required:
    if term not in all_text:
        issues.append(f"缺少关键内容: {term}")

bad_widths = []
empty_cells = []
for ti, t in enumerate(doc.tables, 1):
    widths = [int(x.get(qn("w:w"))) for x in t._tbl.tblGrid if x.get(qn("w:w"))]
    if widths and sum(widths) != 9360:
        bad_widths.append((ti, sum(widths)))
    for ri, row in enumerate(t.rows, 1):
        for ci, cell in enumerate(row.cells, 1):
            if not cell.text.strip():
                empty_cells.append((ti, ri, ci))
if bad_widths:
    issues.append(f"表格宽度异常: {bad_widths}")
if empty_cells:
    issues.append(f"空单元格: {empty_cells[:10]}")

headings = [x.text for x in doc.paragraphs if x.style and x.style.name.startswith("Heading")]
if len(headings) < 25:
    issues.append(f"标题层级不足: {len(headings)}")
if not doc.sections[0].header.paragraphs[0].text.strip():
    issues.append("页眉为空")

print({
    "file_size": p.stat().st_size,
    "paragraphs": len(doc.paragraphs),
    "headings": len(headings),
    "tables": len(doc.tables),
    "issues": issues,
})
raise SystemExit(1 if issues else 0)
