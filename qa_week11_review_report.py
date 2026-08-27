from pathlib import Path
from zipfile import ZipFile, BadZipFile
from docx import Document
from docx.oxml.ns import qn

p = Path(r"D:\PycharmProjects\PythonProject6\outputs\week11_simulated_dataset\模拟数据复盘分析报告.docx")
issues = []
try:
    with ZipFile(p) as z:
        bad = z.testzip()
        if bad:
            issues.append(f"损坏的ZIP成员: {bad}")
except BadZipFile:
    issues.append("DOCX不是有效ZIP包")

doc = Document(p)
text = "\n".join(x.text for x in doc.paragraphs)
required = ["模拟数据，不代表百度真实业务表现", "90.1%", "81.2%", "86.9%", "68.4%", "任务自动化", "管理人员", "产品增长"]
for item in required:
    if item not in text and not any(item in c.text for t in doc.tables for r in t.rows for c in r.cells):
        issues.append(f"缺少关键内容: {item}")

empty_cells = []
bad_widths = []
for ti, t in enumerate(doc.tables, 1):
    for ri, row in enumerate(t.rows, 1):
        for ci, cell in enumerate(row.cells, 1):
            if not cell.text.strip():
                empty_cells.append((ti, ri, ci))
    grid = t._tbl.tblGrid
    widths = [int(c.get(qn("w:w"))) for c in grid if c.get(qn("w:w"))]
    if widths and sum(widths) != 9360:
        bad_widths.append((ti, sum(widths)))

headings = [x.text for x in doc.paragraphs if x.style and x.style.name.startswith("Heading")]
if len(headings) < 20:
    issues.append(f"标题数量偏少: {len(headings)}")
if empty_cells:
    issues.append(f"存在空单元格: {empty_cells[:10]}")
if bad_widths:
    issues.append(f"表格宽度异常: {bad_widths}")
if not doc.sections[0].header.paragraphs[0].text.strip():
    issues.append("页眉为空")

print({
    "file_size": p.stat().st_size,
    "paragraphs": len(doc.paragraphs),
    "headings": len(headings),
    "tables": len(doc.tables),
    "sections": len(doc.sections),
    "issues": issues,
})
raise SystemExit(1 if issues else 0)
