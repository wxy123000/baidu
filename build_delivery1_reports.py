from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "交付1-完整报告"
FIG = ROOT / ".generated_report_figures"
BLUE = "5D6DF4"
BLUE_DARK = "2E4A8A"
BLUE_LIGHT = "EEF2FF"
GRAY = "667085"
LIGHT = "F2F4F7"
GREEN = "2E9D68"
FONT_CN = "微软雅黑"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths_dxa, indent=120):
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[idx]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    paragraph.add_run(" 页")


def set_run_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = FONT_CN
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CN)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if italic is not None:
        run.italic = italic


def setup_doc(title, subtitle):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.82)
    sec.bottom_margin = Inches(0.78)
    sec.left_margin = Inches(1.0)
    sec.right_margin = Inches(1.0)
    sec.header_distance = Inches(0.35)
    sec.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_CN
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE_DARK, 16, 8),
        ("Heading 2", 13, BLUE_DARK, 12, 6),
        ("Heading 3", 11.5, "1F4D78", 8, 4),
    ):
        st = styles[name]
        st.font.name = FONT_CN
        st._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
    for name in ("List Number", "List Bullet"):
        st = styles[name]
        st.font.name = FONT_CN
        st._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
        st.font.size = Pt(10.5)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.line_spacing = 1.25

    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hr = hp.add_run("百度办公 AI 助手｜第10周交付材料")
    set_run_font(hr, 9, False, GRAY)
    add_page_field(sec.footer.paragraphs[0])
    for r in sec.footer.paragraphs[0].runs:
        set_run_font(r, 9, False, GRAY)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    set_run_font(r, 23, True, BLUE_DARK)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(14)
    r = p2.add_run(subtitle)
    set_run_font(r, 11.5, False, GRAY)
    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(16)
    r = meta.add_run("项目：百度办公 AI 助手 MVP　｜　阶段：MVP测试与优化　｜　日期：2026年8月")
    set_run_font(r, 9.5, False, GRAY)
    return doc


def add_heading(doc, text, level=1):
    return doc.add_paragraph(text, style=f"Heading {level}")


def add_para(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        a, b = text[:len(bold_lead)], text[len(bold_lead):]
        set_run_font(p.add_run(a), 10.5, True, BLUE_DARK)
        set_run_font(p.add_run(b), 10.5)
    else:
        set_run_font(p.add_run(text), 10.5)
    return p


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        set_run_font(p.add_run(item), 10.5)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        set_run_font(p.add_run(item), 10.5)


def add_callout(doc, label, text, fill=BLUE_LIGHT):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_run_font(p.add_run(label + "　"), 10.5, True, BLUE_DARK)
    set_run_font(p.add_run(text), 10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_table(doc, headers, rows, widths, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, LIGHT)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(header), font_size, True, BLUE_DARK)
    set_repeat_table_header(table.rows[0])
    for row_data in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_data):
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_run_font(p.add_run(str(value)), font_size)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_figure(doc, path, caption, note=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(6.25))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(3)
    cap.paragraph_format.space_after = Pt(2)
    set_run_font(cap.add_run(caption), 9.5, True, BLUE_DARK)
    if note:
        n = doc.add_paragraph()
        n.alignment = WD_ALIGN_PARAGRAPH.CENTER
        n.paragraph_format.space_after = Pt(6)
        set_run_font(n.add_run(note), 8.5, False, GRAY)


def build_dashboard_report():
    doc = setup_doc("项目核心数据看板原型设计", "百度办公 AI 助手 MVP｜核心指标、页面结构与可视化方案")
    add_callout(doc, "报告定位", "本报告聚焦数据看板原型及指标展示设计。图中数值均为原型展示数据，不代表系统实际运行结果。")

    add_heading(doc, "一、项目背景与设计目标")
    add_para(doc, "随着百度办公AI助手MVP完成基础开发，系统已经具备资料上传与文本输入、会议纪要生成、文档问答、任务拆解、周报生成、自然语言咨询、智能推荐、任务自动化、结果导出、用户反馈和数据指标统计等功能。产品进入测试阶段后，仅判断页面能否运行已不足以说明MVP是否真正可用，还需要通过统一的数据指标观察用户是否进入系统、是否完成核心操作、AI调用是否稳定以及用户是否认可生成结果。")
    add_para(doc, "本报告围绕百度办公AI助手MVP的数据监测需求，设计项目核心数据看板原型，明确数据来源、采集事件、指标口径、页面模块和展示方式，为功能完整性测试、AI能力专项测试及后续产品优化提供数据基础。")
    add_para(doc, "数据看板的建设目标包括以下六点：")
    add_numbered(doc, [
        "判断MVP是否被用户真实访问和使用；",
        "判断用户能否完成资料输入、功能选择、AI生成、导出和反馈等核心流程；",
        "比较不同办公功能的使用频率、完成情况和用户价值；",
        "评估千帆模型调用、本地降级和结果返回过程是否稳定；",
        "分析用户评分、帮助程度和再次使用意愿；",
        "及时发现上传、生成、任务执行、导出和反馈过程中的异常。",
    ])

    add_heading(doc, "二、数据看板需求分析")
    add_para(doc, "MVP阶段的数据看板不追求复杂的数据仓库或商业智能平台，而是以“快速判断、及时预警、支持复盘”为核心。看板应当能够在有限测试样本下呈现关键流程结果，并避免使用含义不清或无法采集的指标。")
    add_heading(doc, "2.1 使用规模与核心流程需求", 2)
    add_para(doc, "看板首先需要呈现访问会话数、生成请求数、成功生成次数和反馈次数，判断产品是否获得真实使用。同时通过上传成功、功能触发、结果返回、导出保存和反馈提交等关键事件，观察用户是否能够完成完整任务链路。")
    add_heading(doc, "2.2 功能价值分析需求", 2)
    add_para(doc, "系统包含会议纪要、文档问答、任务拆解、周报生成、自然语言咨询、智能推荐和任务自动化等功能。看板需要在统一统计周期内比较各功能的使用次数、成功率、任务完成率和用户反馈，识别高频功能、低完成率环节及后续优先优化方向。")
    add_heading(doc, "2.3 AI稳定性与用户体验需求", 2)
    add_para(doc, "AI能力分析既要观察千帆接口是否正常，也要区分真实模型结果与本地降级结果。除成功率外，还应统计总体结果返回率、API调用失败率和平均生成时间。用户体验方面，应结合评分、帮助率、再次使用意愿和文字反馈综合判断，避免仅依据单一指标作出结论。")

    add_heading(doc, "三、数据来源与指标体系")
    add_heading(doc, "3.1 数据来源", 2)
    add_para(doc, "数据看板主要使用系统运行过程中产生的匿名事件日志、AI调用结果、用户反馈和异常信息。系统数据指标日志不保存用户上传文档正文，用户标识和会话标识采用随机编号，以降低敏感信息泄露风险。")
    add_table(doc, ["数据来源", "主要记录内容", "主要用途"], [
        ["系统事件日志", "访问、上传、功能触发、生成、导出和任务执行事件", "统计使用规模、功能使用及核心流程完成情况"],
        ["AI调用记录", "调用方式、成功状态、处理耗时、失败原因和降级情况", "评估千帆接口及AI生成稳定性"],
        ["用户反馈记录", "评分、帮助程度、再次使用意愿和文字意见", "评估生成结果质量与用户满意度"],
        ["系统异常信息", "上传、生成、导出、任务执行和反馈提交异常", "发现Bug并支持风险排查"],
    ], [1800, 3900, 3660])

    add_heading(doc, "3.2 核心事件设计", 2)
    add_table(doc, ["事件名称", "触发条件", "对应指标"], [
        ["session_started", "用户进入系统并形成新会话", "访问会话数"],
        ["upload_success / failed", "资料解析成功或失败", "上传成功率、上传失败率"],
        ["generation_requested", "用户发起AI生成请求", "生成请求数"],
        ["qianfan_succeeded / failed", "千帆接口成功或失败", "千帆生成成功率、API调用失败率"],
        ["local_fallback_used", "接口失败后使用本地演示结果", "本地降级率"],
        ["result_displayed", "页面成功展示生成结果", "总体结果返回率"],
        ["export_success / failed", "结果导出成功或失败", "导出成功率、导出失败率"],
        ["task_completed", "用户完成一次核心办公任务", "任务完成率"],
        ["feedback_submitted", "用户提交评分或文字反馈", "评分覆盖率、帮助率"],
    ], [2500, 4200, 2660])

    add_heading(doc, "3.3 核心指标口径", 2)
    add_table(doc, ["指标名称", "计算口径", "分析作用"], [
        ["访问会话数", "统计周期内去重后的session_id数量", "判断MVP是否被实际访问"],
        ["生成请求数", "generation_requested事件总次数", "判断AI功能使用规模"],
        ["功能使用次数", "对应功能被正式触发的次数", "比较不同功能的使用频率"],
        ["千帆生成成功率", "千帆成功返回次数 ÷ 生成请求总次数 × 100%", "判断真实模型调用稳定性"],
        ["本地降级率", "本地降级次数 ÷ 生成请求总次数 × 100%", "判断接口失败后的降级情况"],
        ["总体结果返回率", "页面成功展示结果次数 ÷ 生成请求总次数 × 100%", "判断用户是否最终获得结果"],
        ["任务完成率", "完成核心任务的会话数 ÷ 启动核心任务的会话数 × 100%", "判断核心流程是否顺畅"],
        ["平均生成时间", "成功生成任务总耗时 ÷ 成功生成次数", "判断等待时间是否可接受"],
        ["用户平均评分", "有效评分总和 ÷ 有效评分数量", "判断用户满意程度"],
        ["评分覆盖率", "提交评分的任务数 ÷ 成功生成任务数 × 100%", "判断反馈样本是否充足"],
        ["帮助率", "选择“有帮助”的反馈数 ÷ 有效反馈数 × 100%", "判断结果是否具有实际价值"],
        ["再次使用意愿率", "选择“愿意再次使用”的反馈数 ÷ 有效反馈数 × 100%", "判断产品后续使用潜力"],
    ], [1700, 4750, 2910], 8.7)

    add_heading(doc, "3.4 统计周期与更新方式", 2)
    add_para(doc, "看板支持按日、按周和MVP测试阶段查看数据。每日数据用于检查接口和系统异常，每周数据用于比较功能表现与用户反馈，阶段数据用于判断功能优先级和验收结果。MVP阶段可在进入数据指标页面或刷新页面时重新汇总SQLite数据；当数据量增加后，可进一步采用定时聚合和缓存机制。")

    add_heading(doc, "四、数据看板整体结构")
    add_para(doc, "根据百度办公AI助手MVP的核心业务流程、AI能力测试要求和数据复盘需求，数据看板划分为用户概览、功能使用、AI能力效果、新增功能、用户反馈、异常与风险六个模块。")
    add_table(doc, ["模块名称", "展示内容", "设计目的"], [
        ["用户概览模块", "访问会话数、事件次数、生成次数、反馈次数", "判断MVP是否被真实使用"],
        ["功能使用模块", "各功能使用次数、核心功能使用率、任务完成率", "判断哪些功能更有价值"],
        ["AI能力效果模块", "千帆生成成功率、本地降级率、总体结果返回率、API调用失败率", "判断AI调用是否稳定"],
        ["新增功能模块", "咨询成功率、推荐采用率、Function Calling成功率、任务状态更新率", "判断新版功能表现"],
        ["用户反馈模块", "用户评分、帮助率、再次使用意愿率、有效反馈数量", "判断用户满意度"],
        ["异常与风险模块", "上传、生成、导出、任务和反馈失败情况", "帮助发现Bug和运行风险"],
    ], [1800, 4800, 2760])
    add_para(doc, "六个模块分别覆盖产品使用规模、功能价值、AI稳定性、新版功能表现、用户满意度和系统风险，共同构成MVP测试阶段的数据监测体系。")

    add_heading(doc, "五、数据看板页面原型设计")
    add_heading(doc, "5.1 页面整体布局", 2)
    add_para(doc, "看板采用“顶部核心指标＋中部趋势分析＋底部反馈与异常明细”的结构。顶部用于快速浏览关键结果，中部用于比较功能使用和AI调用趋势，底部用于查看用户反馈、失败记录和待处理问题。页面提供统计周期筛选，使测试人员能够在日、周和阶段视角之间切换。")

    add_heading(doc, "5.2 顶部核心指标区", 2)
    add_para(doc, "顶部使用六张数字卡片展示访问会话数、生成请求数、千帆生成成功率、本地降级率、任务完成率和用户平均评分。数字卡片用于快速判断MVP的使用规模、AI稳定性和整体体验。")
    add_figure(doc, FIG / "图1-数据看板顶部核心指标区.png", "图1　数据看板顶部核心指标区原型图", "注：图中数据为原型展示数据，不代表系统实际运行结果。")

    add_heading(doc, "5.3 功能使用分析区", 2)
    add_para(doc, "功能使用区域以横向柱状图比较会议纪要、文档问答、任务拆解、周报生成、自然语言咨询、智能推荐和任务自动化的使用次数，并配合任务完成率、推荐采用率等摘要指标，帮助团队识别高频功能和低使用功能。")
    add_table(doc, ["功能名称", "主要观察指标"], [
        ["会议纪要生成", "使用次数、生成成功率、导出率"],
        ["文档问答", "提问次数、回答成功率、依据有效性反馈"],
        ["任务拆解", "使用次数、任务完成率、保存率"],
        ["周报生成", "使用次数、导出率、用户评分"],
        ["自然语言咨询", "咨询请求次数、咨询成功率、多轮对话率"],
        ["智能推荐", "推荐触发次数、推荐采用率、推荐修改率"],
        ["任务自动化", "Function Calling成功率、结构化任务有效率、状态更新率"],
    ], [2400, 6960])
    add_figure(doc, FIG / "图2-功能使用情况分析.png", "图2　功能使用情况分析原型图", "注：图中数据为原型展示数据，用于说明图表布局和比较方式。")

    add_heading(doc, "5.4 AI能力效果分析区", 2)
    add_para(doc, "AI能力效果区域重点区分千帆真实模型结果与本地降级结果，避免将页面成功返回误认为千帆接口稳定。建议使用百分比卡片展示成功率和降级率，使用折线图展示不同日期的生成成功率和平均耗时变化，并提供失败原因明细。")
    add_table(doc, ["指标名称", "展示方式", "判断重点"], [
        ["千帆生成成功率", "百分比卡片＋趋势线", "真实模型调用是否稳定"],
        ["本地降级率", "百分比卡片＋趋势线", "降级是否过于频繁"],
        ["总体结果返回率", "百分比卡片", "用户是否最终看到结果"],
        ["API调用失败率", "趋势图＋失败原因", "接口异常是否集中发生"],
        ["平均生成时间", "折线图", "用户等待时间是否可接受"],
    ], [2100, 2700, 4560])

    add_heading(doc, "5.5 用户反馈分析区", 2)
    add_para(doc, "用户反馈区域展示平均评分、帮助率、再次使用意愿率、评分覆盖率和有效文字反馈数。定量指标用于观察整体趋势，文字反馈用于解释低评分原因。分析时应区分AI内容质量、页面操作体验、生成速度和功能缺失等不同问题类型。")

    add_heading(doc, "5.6 异常与风险监控区", 2)
    add_para(doc, "异常与风险区域用于集中展示上传失败、生成失败、导出失败、Function Calling失败、字段缺失和反馈提交失败等情况。每类异常应提供发生时间、关联功能、错误类型和处理状态，以便测试人员快速定位问题。")

    add_heading(doc, "六、原型与系统实现说明")
    add_para(doc, "本报告中的数据看板图片属于MVP阶段的页面原型，主要用于说明信息结构、指标布局和可视化方式，图中数值不作为项目真实测试结论。当前系统已经具备SQLite事件记录、匿名会话标识、AI调用结果记录、用户反馈保存和数据指标查看等基础能力，可为核心指标提供数据来源。")
    add_para(doc, "后续实现应优先保证指标采集和计算口径正确，再逐步完善趋势图、筛选器、异常明细和问题状态等可视化能力。看板不应为了追求展示效果而虚构无法由系统采集的数据；暂未完成采集的指标应标记为“待接入”或从当前版本中移除。")
    add_callout(doc, "实现边界", "系统无需完全照搬原型图片的视觉细节，但应能够采集或计算报告中声明的核心指标，并明确区分示例数据与实际测试数据。", "EAF7EF")

    add_heading(doc, "七、总结")
    add_para(doc, "本报告完成了百度办公AI助手MVP核心数据看板的原型设计，明确了数据来源、核心事件、指标口径、六大模块和页面展示方式。看板覆盖用户使用、功能表现、AI能力、用户反馈及异常风险等关键维度，可以帮助团队判断产品是否被真实使用、核心流程是否顺畅以及千帆模型调用是否稳定。")
    add_para(doc, "数据看板是功能测试和AI能力专项测试的数据基础。后续应结合真实测试数据持续校验指标口径，并与《数据复盘机制规范》配合使用，将异常指标和用户反馈转化为明确的优化任务。")
    return doc


def build_review_report():
    doc = setup_doc("数据复盘机制规范", "百度办公 AI 助手 MVP｜复盘周期、预警规则与问题闭环")
    add_callout(doc, "规范定位", "本规范以《项目核心数据看板原型设计》中确定的数据指标为基础，规定数据检查、分析、预警和闭环处理的方法。")

    add_heading(doc, "一、规范背景与目的")
    add_para(doc, "数据看板能够集中展示产品使用、AI调用、用户反馈和系统异常情况，但仅展示数据不能自动推动产品改进。如果缺少统一的复盘周期、指标阈值、职责分工和问题闭环要求，团队可能出现指标口径不一致、异常发现不及时、问题无人负责或修复后未复测等情况。")
    add_para(doc, "本规范适用于百度办公AI助手MVP测试阶段的数据检查与产品复盘，主要目标是统一复盘输入和分析方法，及时识别异常变化，将数据问题转化为可执行的优化任务，并通过修复、复测和结果记录形成持续改进闭环。")

    add_heading(doc, "二、复盘范围与基本原则")
    add_heading(doc, "2.1 复盘范围", 2)
    add_para(doc, "复盘范围覆盖用户使用、功能表现、AI能力、新增功能、用户反馈和异常风险六个维度。复盘使用的数据包括匿名事件日志、AI调用结果、用户评分与文字反馈、系统异常信息以及上一周期的问题处理记录。")
    add_heading(doc, "2.2 基本原则", 2)
    add_numbered(doc, [
        "口径统一：所有参与人员使用同一指标定义、统计周期和计算方式。",
        "异常优先：影响核心流程、数据安全或大范围用户使用的问题优先处理。",
        "数据与反馈结合：定量指标用于发现变化，用户反馈用于解释变化原因。",
        "责任明确：每个问题必须确定负责人、优先级和计划完成时间。",
        "修复必须复测：问题完成修复后必须按照原场景再次测试并记录结果。",
        "最小必要采集：复盘不读取或传播用户上传文档正文等敏感信息。",
    ])

    add_heading(doc, "三、复盘周期与内容")
    add_para(doc, "百度办公AI助手MVP采用“每日检查、每周复盘、阶段总结”三级机制。不同周期关注的问题和输出结果不同，避免每日检查过度分析，也避免阶段总结遗漏日常异常。")
    add_table(doc, ["复盘周期", "复盘内容", "负责人", "主要输出"], [
        ["每日检查", "检查数据是否正常记录，是否出现接口、上传、生成、导出或反馈异常", "测试人员／数据负责人", "每日异常记录"],
        ["每周复盘", "分析功能使用、任务完成、AI稳定性、用户评分和文字反馈", "产品负责人牵头，开发与测试参与", "周度指标简报、优化清单"],
        ["阶段总结", "判断MVP验收结果、功能优先级、Prompt和页面流程优化方向", "项目负责人", "阶段复盘报告、下一阶段计划"],
    ], [1300, 3900, 2260, 1900], 8.8)

    add_heading(doc, "3.1 每日检查", 2)
    add_para(doc, "每日检查以系统稳定性和数据完整性为主，建议在测试结束后或固定时间执行。检查内容包括当日是否产生事件数据、千帆接口是否连续失败、本地降级率是否突然上升、上传和导出是否可用、反馈是否成功写入以及是否出现集中性错误。")
    add_bullets(doc, [
        "确认访问、上传、生成、导出和反馈事件能够正常写入；",
        "检查千帆调用失败、本地降级和平均生成时间变化；",
        "检查上传、生成、导出、任务执行和反馈提交异常；",
        "将影响核心流程的问题立即登记并通知负责人。",
    ])

    add_heading(doc, "3.2 每周复盘", 2)
    add_para(doc, "每周复盘以功能价值和用户体验为主。复盘人员应比较本周与上周的数据变化，重点分析高频功能、低完成率功能、千帆成功率、用户评分和文字反馈。对于明显变化，应结合测试场景和异常日志查找原因，不应仅凭单一百分比直接作出结论。")

    add_heading(doc, "3.3 阶段总结", 2)
    add_para(doc, "阶段总结在一轮MVP测试或重要版本结束后进行。总结应回答产品是否达到阶段验收目标、哪些功能值得继续投入、哪些功能需要调整优先级、Prompt是否需要优化、页面流程是否存在明显阻塞以及上一阶段的问题是否已经闭环。")

    add_heading(doc, "四、参与人员与职责")
    add_table(doc, ["角色", "主要职责", "交付要求"], [
        ["测试人员", "执行测试、检查异常数据、登记Bug并完成复测", "测试记录和复测结论完整"],
        ["数据负责人", "汇总指标、检查数据口径和异常变化", "指标数据可追溯、计算一致"],
        ["产品负责人", "分析功能表现与用户反馈，提出优化方案", "形成明确的优化优先级"],
        ["开发人员", "排查接口和系统问题，完成修复并说明影响范围", "提交修复结果和必要说明"],
        ["项目负责人", "确定处理优先级、协调资源并确认阶段结论", "检查问题闭环和阶段目标"],
    ], [1500, 4700, 3160])
    add_para(doc, "MVP阶段部分角色可由同一项目成员兼任，但每项职责仍应分别执行，避免因人员较少而省略数据核对、问题登记或复测环节。")

    add_heading(doc, "五、异常预警规则")
    add_para(doc, "预警阈值用于帮助团队快速识别需要关注的变化，不直接等同于最终验收结论。MVP初期样本量较小时，应同时查看绝对次数和百分比，避免因少量样本造成误判。")
    add_table(doc, ["指标", "初始预警条件", "处理方向"], [
        ["千帆生成成功率", "低于90%", "检查API权限、模型配置、网络和超时设置"],
        ["本地降级率", "高于10%", "分析千帆调用失败原因及降级触发逻辑"],
        ["总体结果返回率", "低于95%", "检查后端生成、异常处理和页面展示逻辑"],
        ["任务完成率", "低于70%", "检查功能入口、页面引导和操作步骤"],
        ["用户平均评分", "低于3.5分", "分类分析内容质量、速度和页面体验反馈"],
        ["导出失败率", "高于5%", "检查Word、PDF和TXT导出逻辑"],
        ["平均生成时间", "高于15秒", "检查请求参数、模型响应和超时策略"],
        ["反馈提交失败率", "高于2%", "检查表单校验和SQLite写入"],
    ], [2200, 2200, 4960])
    add_callout(doc, "阈值说明", "以上数值为MVP测试阶段的初始参考标准。后续应根据真实样本量、业务目标和历史基线调整，并在每次调整时记录原因和生效时间。")

    add_heading(doc, "六、数据复盘流程")
    add_para(doc, "数据复盘按照“数据采集—看板展示—指标分析—发现问题—提出优化方案—修复或调整—再次测试”的顺序执行。再次测试结果重新进入数据采集环节，形成持续迭代闭环。")
    add_figure(doc, FIG / "图3-数据复盘流程图.png", "图1　数据复盘闭环流程图")
    add_heading(doc, "6.1 数据准备", 2)
    add_para(doc, "复盘前确认统计周期、指标口径和数据完整性。若出现日志缺失、重复记录或样本量过小，应先标记数据限制，再进行分析。")
    add_heading(doc, "6.2 问题识别", 2)
    add_para(doc, "将指标结果与预警阈值、上一周期结果和测试目标进行比较。发现异常后，结合事件明细、失败原因和用户文字反馈判断问题发生在哪个功能或流程环节。")
    add_heading(doc, "6.3 优化与复测", 2)
    add_para(doc, "优化方案应明确修改对象、预期结果、负责人和完成时间。修复完成后使用相同测试场景再次验证，并比较修复前后的指标、错误记录和用户体验变化。")

    add_heading(doc, "七、问题分级与闭环管理")
    add_heading(doc, "7.1 问题分级", 2)
    add_table(doc, ["级别", "判定标准", "处理要求"], [
        ["高", "核心流程不可用、数据安全风险、大范围生成失败或持续接口异常", "立即处理，优先于一般优化任务"],
        ["中", "部分功能受影响、完成率明显下降或用户体验受到较大影响", "纳入当前迭代并安排复测"],
        ["低", "轻微界面问题、个别文案问题或不影响主流程的改进项", "进入优化清单，按优先级处理"],
    ], [1300, 5100, 2960])

    add_heading(doc, "7.2 问题记录字段", 2)
    add_para(doc, "每个问题至少记录问题编号、发现日期、数据依据、关联功能、问题描述、严重程度、负责人、计划完成时间、处理状态和复测结果。状态统一使用“待确认、待处理、处理中、待复测、已关闭”五种。")
    add_table(doc, ["问题编号", "关联指标／功能", "问题描述", "优先级", "负责人", "状态"], [
        ["BUG-001", "本地降级率", "千帆接口异常导致降级率连续升高", "高", "开发人员", "处理中"],
        ["UX-001", "任务完成率", "用户未能快速发现下一步操作入口", "中", "产品负责人", "待处理"],
        ["AI-001", "用户评分", "周报生成内容结构与输入信息匹配度不足", "中", "产品／开发", "待确认"],
    ], [1300, 1800, 2800, 1000, 1300, 1160], 8.5)

    add_heading(doc, "7.3 关闭条件", 2)
    add_para(doc, "问题只有在修复内容已提交、原测试场景复测通过、未引入新的严重异常且处理结果完成记录后，才能标记为“已关闭”。若复测仍失败，应重新打开问题并更新原因分析和处理计划。")

    add_heading(doc, "八、复盘输出与记录要求")
    add_para(doc, "每次复盘应形成结构化记录，确保下一周期能够追踪变化和核对处理结果。")
    add_bullets(doc, [
        "核心指标简报：记录统计周期、样本量、关键结果和环比变化；",
        "异常问题清单：记录触发预警的指标、影响范围和初步原因；",
        "用户反馈摘要：归纳高频意见、典型问题和正向反馈；",
        "优化任务清单：明确任务内容、优先级、负责人和完成时间；",
        "上期问题复测结果：说明问题是否解决、是否关闭以及遗留风险。",
    ])
    add_para(doc, "复盘记录应存放在统一项目目录中，文件名建议包含日期和复盘周期，例如“2026-08-09_每周数据复盘记录”。涉及测试用户反馈时，不应在复盘材料中复制敏感信息或完整文档正文。")

    add_heading(doc, "九、典型指标处理示例")
    add_table(doc, ["数据表现", "可能原因", "建议处理"], [
        ["本地降级率较高", "接口权限、模型配置、网络、超时或异常处理存在问题", "检查失败日志和千帆配置，复测主要生成功能"],
        ["推荐采用率较低", "推荐规则不准确、推荐理由不清晰或入口不明显", "优化推荐规则、提示词及页面展示"],
        ["任务完成率较低", "用户在上传、选择功能、确认或导出环节遇到阻塞", "分析关键事件漏斗并检查页面引导"],
        ["用户评分较低", "AI内容质量、生成速度或页面体验未达到预期", "结合文字反馈分类处理，避免仅调整Prompt"],
    ], [1900, 3900, 3560], 8.8)

    add_heading(doc, "十、总结")
    add_para(doc, "本规范建立了百度办公AI助手MVP的每日检查、每周复盘和阶段总结机制，明确了复盘原则、参与角色、初始预警阈值、问题分级、闭环流程和输出要求。通过统一指标口径和问题记录方式，团队可以及时识别AI调用异常、核心流程阻塞和用户体验问题。")
    add_para(doc, "数据复盘的最终目的不是简单汇报数字，而是将数据变化转化为明确的优化行动。后续应根据实际测试样本持续调整预警阈值，并坚持“发现问题—责任分配—修复调整—再次测试—结果记录”的闭环要求，使百度办公AI助手逐步从可运行原型发展为可测试、可评估和可持续优化的AI办公产品。")
    return doc


def structural_audit(path):
    doc = Document(path)
    assert len(doc.paragraphs) > 25
    assert any(p.style.name.startswith("Heading") for p in doc.paragraphs)
    assert len(doc.tables) >= 5
    for table in doc.tables:
        assert table.rows and table.columns
        for row in table.rows:
            assert len(row.cells) == len(table.columns)
    return {
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "inline_shapes": len(doc.inline_shapes),
        "sections": len(doc.sections),
    }


def main():
    OUT.mkdir(exist_ok=True)
    dashboard_path = OUT / "项目核心数据看板原型设计-完整版.docx"
    review_path = OUT / "数据复盘机制规范-完整版.docx"
    build_dashboard_report().save(dashboard_path)
    build_review_report().save(review_path)
    print(dashboard_path)
    print(structural_audit(dashboard_path))
    print(review_path)
    print(structural_audit(review_path))


if __name__ == "__main__":
    main()
