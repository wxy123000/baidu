from docx import Document

p = r"D:\PycharmProjects\PythonProject6\第十周项目周报.docx"
d = Document(p)
print("paragraphs", len(d.paragraphs), "tables", len(d.tables), "sections", len(d.sections))
for i, x in enumerate(d.paragraphs):
    t = " ".join(x.text.split())
    if t:
        print(i, x.style.name, t[:180])
for i, t in enumerate(d.tables):
    print("TABLE", i + 1, len(t.rows), len(t.columns), [c.text for c in t.rows[0].cells])
