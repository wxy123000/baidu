from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path("AC-010长文档分段处理测试样本.docx")


def set_run_font(run, name="等线", size=11, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color="6B7280")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])
    tail = paragraph.add_run(" 页")
    set_run_font(tail, size=9, color="6B7280")


def add_marker(doc, title, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.25
    r1 = p.add_run(f"{title}：")
    set_run_font(r1, size=12, bold=True, color="C00000")
    r2 = p.add_run(value)
    set_run_font(r2, size=12, bold=True, color="C00000")


def add_filler(doc, start_index, count, phase):
    base = (
        "本段为长文档分段处理验收使用的普通项目说明文字，用于模拟真实项目资料中的需求背景、"
        "业务流程、执行安排、协作机制、质量标准、风险跟踪和阶段复盘。测试时不需要概括本段，"
        "只需要根据文档指定位置的唯一识别信息回答问题。为确保测试结果可复现，普通段落不会出现"
        "任何测试识别码，也不会重复对应的人员、金额或发布日期。"
    )
    for i in range(start_index, start_index + count):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.25
        run = p.add_run(f"{phase}阶段资料第{i:03d}段。{base}")
        set_run_font(run)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "等线"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(12)
    run = title.add_run("AC-010 长文档分段处理测试样本")
    set_run_font(run, size=18, bold=True, color="2E74B5")

    note = doc.add_paragraph()
    note.paragraph_format.space_after = Pt(12)
    note.paragraph_format.line_spacing = 1.25
    r = note.add_run(
        "测试说明：将本文件上传至系统并选择“文档问答”，分别询问开头、中间和结尾的识别信息。"
        "三处均回答正确才表示完整通过长文档分段处理验收。"
    )
    set_run_font(r, bold=True, color="1F4D78")

    add_marker(doc, "开头识别信息", "项目代号为START-2026，启动负责人为张明。")
    add_filler(doc, 1, 150, "前部")
    add_marker(doc, "中间识别信息", "核心验收编号为MIDDLE-8888，中期预算为12万元。")
    add_filler(doc, 151, 90, "后部")
    add_marker(doc, "结尾识别信息", "最终发布编号为END-9999，发布日期为10月20日。")

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hr = header.add_run("AC-010 长文档测试样本")
    set_run_font(hr, size=9, color="6B7280")
    add_page_number(section.footer.paragraphs[0])

    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build().resolve())
