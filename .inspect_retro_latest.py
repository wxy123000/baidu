from docx import Document
p=r'D:\PycharmProjects\PythonProject6\.doc_review_data_retro_latest.docx'
d=Document(p)
print('PARAGRAPHS',len(d.paragraphs),'TABLES',len(d.tables),'IMAGES',len(d.inline_shapes))
for i,p in enumerate(d.paragraphs):
 t=' '.join(p.text.split())
 if t: print(f'P{i:03d}\t[{p.style.name}]\t{t}')
print('\n=== TABLES ===')
for ti,t in enumerate(d.tables):
 print(f'\nTABLE {ti+1} {len(t.rows)}x{len(t.columns)}')
 for ri,row in enumerate(t.rows): print(f'R{ri:02d}\t'+' | '.join(' '.join(c.text.split()) for c in row.cells))
