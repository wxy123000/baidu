from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUT = Path(r"D:\PycharmProjects\PythonProject6\outputs\week11_simulated_dataset")
SUMMARY = json.loads((OUT / "dataset_summary.json").read_text(encoding="utf-8"))
DOCX = OUT / "模拟用户数据集构建说明.docx"
CHART = OUT / "xlsx_qa" / "汇总看板.png"

NAVY = "1F4E78"
BLUE = "4472C4"
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "667085"
INK = "1F2937"
GREEN = "70AD47"
AMBER = "ED7D31"
RED = "C00000"
WHITE = "FFFFFF"


def set_cell_shading(cell, color: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_width(cell, width_dxa: int):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_fixed(table, widths: list[int]):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        for i, cell in enumerate(row.cells):
            set_cell_width(cell, widths[i])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_font(run, name="Microsoft YaHei", size=10.5, bold=False, color=INK, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def style_para(p, before=0, after=6, line=1.1, align=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if align is not None:
        p.alignment = align


def add_body(doc, text, bold_prefix=None, after=6):
    p = doc.add_paragraph()
    style_para(p, after=after, line=1.15)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_font(r2)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    style_para(p, after=4, line=1.15)
    set_font(p.add_run(text))
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    style_para(p, after=4, line=1.15)
    set_font(p.add_run(text))
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_font(r, size={1:16, 2:13, 3:11.5}[level], bold=True, color=NAVY if level < 3 else "2F5597")
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    style_para(p, before=4, after=4, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_font(p.add_run(text), size=9.5, color=MID_GRAY)
    return p


def add_table(doc, headers, rows, widths, header_fill=BLUE, font_size=9.0):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_fixed(table, widths)
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, header_fill)
        p = cell.paragraphs[0]
        style_para(p, after=0, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_font(p.add_run(str(h)), size=font_size, bold=True, color=WHITE)
    for row_data in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_data):
            p = cells[i].paragraphs[0]
            style_para(p, after=0, line=1.0, align=WD_ALIGN_PARAGRAPH.LEFT if i else WD_ALIGN_PARAGRAPH.CENTER)
            set_font(p.add_run(str(value)), size=font_size, color=INK)
        if len(table.rows) % 2 == 1:
            for c in cells:
                set_cell_shading(c, "F8FAFC")
    set_table_fixed(table, widths)
    return table


def add_callout(doc, title, text, color=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_fixed(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, color)
    p = cell.paragraphs[0]
    style_para(p, after=2, line=1.1)
    set_font(p.add_run(title + "："), bold=True, color=NAVY)
    set_font(p.add_run(text), color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_font(run, size=9, color=MID_GRAY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    run2 = paragraph.add_run(" 页")
    set_font(run2, size=9, color=MID_GRAY)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(.492)
section.footer_distance = Inches(.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.1
for level, size, before, after in [(1,16,16,8),(2,13,12,6),(3,11.5,8,4)]:
    s = styles[f"Heading {level}"]
    s.font.name = "Microsoft YaHei"
    s._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = RGBColor.from_string(NAVY if level < 3 else "2F5597")
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True

# Header and footer
hp = section.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
set_font(hp.add_run("百度办公 AI 助手 MVP｜第11周交付材料"), size=9, color=MID_GRAY)
fp = section.footer.paragraphs[0]
add_page_number(fp)

# Cover (memo masthead pattern)
p = doc.add_paragraph()
style_para(p, before=36, after=6)
set_font(p.add_run("WEEK 11 DELIVERABLE 01"), size=10.5, bold=True, color=BLUE)
p = doc.add_paragraph()
style_para(p, after=10)
set_font(p.add_run("模拟用户数据集构建说明"), size=26, bold=True, color=NAVY)
p = doc.add_paragraph()
style_para(p, after=30)
set_font(p.add_run("百度办公 AI 助手 MVP｜模拟数据复盘与方案优化迭代"), size=13, color=MID_GRAY)
add_table(doc, ["项目项", "内容"], [
    ["交付物名称", "《模拟用户数据集构建说明》"],
    ["配套数据文件", "《模拟用户数据集.xlsx》"],
    ["模拟周期", SUMMARY["simulatedPeriod"]],
    ["数据规模", f"{SUMMARY['users']}名模拟用户、{SUMMARY['sessions']}次会话、{SUMMARY['events']}条使用事件"],
    ["数据性质", "完全模拟，仅用于内部分析、复盘演练与MVP优化，不代表真实百度运营数据"],
    ["编制日期", "2026年8月15日"],
], [1800,7560], header_fill=NAVY, font_size=10)
doc.add_paragraph()
add_callout(doc, "重要声明", "本文和配套Excel中的用户、行为、反馈、异常及指标均由规则生成，不包含真实姓名、手机号、身份证号、地址或用户文档正文，也不得用于对外经营披露。", "FFF2CC")
doc.add_page_break()

add_heading(doc, "摘要", 1)
add_body(doc, f"为支持第11周“模拟数据复盘与方案优化迭代”，本报告围绕百度办公 AI 助手 MVP 构建一套可复现、可关联、可计算的模拟用户数据集。数据覆盖职场新人、普通职员、项目负责人和管理人员四类模拟用户，覆盖会议纪要、文档问答、任务拆解、周报生成、自然语言咨询、智能推荐和任务自动化七项功能。")
add_body(doc, f"本次以固定随机种子生成{SUMMARY['users']}名模拟用户、{SUMMARY['sessions']}次会话、{SUMMARY['events']}条使用事件、{SUMMARY['feedback']}条反馈和{SUMMARY['anomalies']}条异常记录。所有汇总指标均从原始记录计算，既可用于按用户群体和使用场景开展数据复盘，也可为产品功能、AI融合、用户体验和产品增长四个方向的优化提供输入。")
add_body(doc, "数据构造强调三项边界：第一，模拟参数仅用于形成合理的数据差异，不等同于行业真实均值；第二，任何分析结论都必须标注“基于模拟数据”；第三，数据集不采集、不保存真实个人敏感信息和办公文档正文。")

add_heading(doc, "目录", 1)
for item in ["一、编制背景与目标","二、构建依据与适用边界","三、数据构建原则","四、模拟用户分群设计","五、数据集结构与关联关系","六、核心字段与指标口径","七、模拟参数与生成规则","八、异常与边界场景设计","九、数据生成流程","十、数据质量校验","十一、数据集结果概览","十二、复盘使用方法","十三、局限性与风险控制","十四、结论","附录A：字段清单","附录B：项目参考材料"]:
    add_bullet(doc, item)
doc.add_page_break()

add_heading(doc, "一、编制背景与目标", 1)
add_heading(doc, "1.1 编制背景", 2)
add_body(doc, "第10周已完成核心数据看板原型、功能完整性测试、AI能力效果测试和竞品对比分析。第11周需要在不接入真实用户数据的前提下，构造一套能够运行既有指标体系的模拟数据，并依据不同用户群体、功能和场景分析产品表现，为后续优化MVP原型提供依据。")
add_heading(doc, "1.2 构建目标", 2)
for x in ["建立可复现的模拟用户、会话、使用事件、反馈和异常数据。","验证第9周数据指标体系与第10周数据看板的字段及计算口径。","形成用户群体、产品功能和异常场景之间的差异，支持分层复盘。","为第11周《模拟数据复盘分析报告》和《项目AI产品全案优化迭代方案》提供统一数据底座。","在不使用真实个人信息和真实业务数据的前提下验证分析流程。"]:
    add_bullet(doc, x)
add_heading(doc, "1.3 交付组成", 2)
add_table(doc,["文件","主要内容","使用方式"],[
    ["模拟用户数据集构建说明.docx","方法、口径、规则、质量校验与边界","评审数据设计是否合理"],
    ["模拟用户数据集.xlsx","原始模拟记录、字段字典和公式汇总看板","筛选、透视、复盘和生成后续报告"],
], [2900,3300,3160])

add_heading(doc, "二、构建依据与适用边界", 1)
add_heading(doc, "2.1 项目内部依据", 2)
add_body(doc, "数据设计主要依据前10周项目材料中的用户画像、需求优先级、PRD验收标准、技术架构、用户体验策略、数据指标体系、风险管控要求，以及第10周功能测试和AI能力测试结果。上述材料用于确定功能范围、字段结构和风险场景，不用于证明模拟参数是真实行业均值。")
add_heading(doc, "2.2 适用范围", 2)
add_table(doc,["可用于","不可用于"],[
    ["验证指标口径和数据看板结构","代表真实百度用户行为或经营业绩"],
    ["演练分群、分功能和分场景复盘","作为真实市场规模、转化率或收入结论"],
    ["比较不同优化方案的相对变化","向外部披露为实测数据或行业权威基准"],
    ["发现字段缺失、统计冲突和异常处理问题","训练或评价真实用户个体"],
], [4680,4680])
add_callout(doc, "统一标识要求", "所有图表、表格和结论中应明确标注“模拟数据”或“基于模拟数据分析”，避免与真实运营数据混淆。", LIGHT_BLUE)

add_heading(doc, "三、数据构建原则", 1)
add_table(doc,["原则","具体要求","落实方式"],[
    ["真实性近似","形成合理但不冒充真实的用户与功能差异","按用户群体、偏好功能、接受度和频率设置分布"],
    ["可复现","相同参数应生成相同结果","固定随机种子20260815并记录模拟周期"],
    ["可关联","不同表能够通过主外键追溯","用户ID、会话ID、事件ID贯穿各表"],
    ["可计算","数据类型和枚举值支持指标计算","评分、状态、响应时间和0/1字段采用结构化值"],
    ["差异化","不同用户群和功能不应呈现完全一致的结果","设置不同偏好权重、成功率和响应时间基线"],
    ["安全合规","不包含真实个人敏感信息和文档正文","全部使用匿名模拟编号及非敏感摘要"],
    ["结果可解释","指标变化能够回到生成规则和原始事件","保留错误类型、功能、输入方式及处理状态"],
], [1500,3550,4310])

add_heading(doc, "四、模拟用户分群设计", 1)
add_body(doc, "结合第三周用户画像和办公AI产品的主要使用角色，将模拟用户划分为四类。用户群体不对应真实个人，仅用于构造不同使用动机与行为强度。")
add_caption(doc, "表1  模拟用户群体设置")
add_table(doc,["用户群体","数量","占比","主要需求","高频功能"],[
    ["职场新人",90,"30%","提升基础办公效率、降低内容整理门槛","会议纪要、周报生成、自然语言咨询"],
    ["普通职员",105,"35%","资料处理、内容生成和日常协作","文档问答、会议纪要、周报生成"],
    ["项目负责人",75,"25%","任务规划、资料决策和团队推进","任务拆解、文档问答、智能推荐"],
    ["管理人员",30,"10%","快速掌握进度、审阅结果和辅助决策","周报生成、智能推荐、任务自动化"],
], [1500,900,900,3150,2910])
add_heading(doc, "4.1 用户属性", 2)
for x in ["AI接受度：高、中、低，用于模拟不同用户对AI结果的接受程度。","使用频率：高频、中频、低频，用于控制活跃天数和事件数量。","偏好功能：每名用户配置一项主要偏好功能，约三分之一事件优先命中该功能。","输入方式：上传文件、粘贴文本或直接输入；自然语言咨询以直接输入为主。","设备类型：以桌面端为主，移动端用于模拟轻量查看和反馈场景。"]:
    add_bullet(doc,x)

add_heading(doc, "五、数据集结构与关联关系", 1)
add_caption(doc, "表2  数据表结构")
add_table(doc,["数据表","记录数","主键","关联字段","主要用途"],[
    ["模拟用户",SUMMARY["users"],"用户ID","—","用户分群和属性分析"],
    ["使用事件",SUMMARY["events"],"事件ID","用户ID、会话ID","功能使用、成功率、耗时和转化分析"],
    ["用户反馈",SUMMARY["feedback"],"反馈ID","事件ID、用户ID","评分、帮助率和再次使用意愿分析"],
    ["异常记录",SUMMARY["anomalies"],"异常ID","事件ID","错误类型、等级和处理状态分析"],
    ["汇总看板","公式结果","—","引用上述原始表","展示核心指标与功能表现"],
    ["字段字典","字段说明","—","对应各数据表","统一字段类型、枚举和统计口径"],
], [1550,900,1200,2300,3410])
add_heading(doc, "5.1 关联逻辑", 2)
add_body(doc, "用户表与使用事件表通过用户ID形成一对多关系；一条使用事件属于一个会话，并可产生零条或一条反馈；失败或高耗时事件可能被抽样进入异常记录。复盘时以使用事件为行为主表，再关联用户属性、反馈与异常信息。")
add_callout(doc, "关系说明", "模拟数据以匿名ID实现关联，不保存真实用户身份、上传文档正文或完整AI对话内容。", LIGHT_BLUE)

add_heading(doc, "六、核心字段与指标口径", 1)
add_heading(doc, "6.1 使用事件核心字段", 2)
add_table(doc,["字段","类型/枚举","统计用途"],[
    ["功能类型","7项产品功能","按功能计算使用量、成功率和完成率"],
    ["生成状态","成功/失败","生成成功率=成功事件数÷全部事件数"],
    ["响应时间(秒)","正数","平均响应时间=全部事件响应时间平均值"],
    ["任务完成","0/1","任务完成率=完成事件数÷成功事件数"],
    ["已导出、已保存","0/1","计算导出率和保存率"],
    ["来源有效性","有效/不足/不适用","评价文档问答依据有效性"],
    ["多轮对话","是/否/不适用","分析自然语言咨询连续对话"],
    ["推荐采纳/修改","是/否/不适用","分析智能推荐采纳和修改行为"],
    ["Function Calling/状态更新","成功/失败/不适用","分析任务自动化链路表现"],
    ["错误类型","文本枚举","定位失败原因和风险聚集功能"],
], [2600,2600,4160])
add_heading(doc, "6.2 反馈字段与口径", 2)
add_table(doc,["指标","计算口径","说明"],[
    ["平均评分","有效反馈评分之和÷有效反馈数","评分范围1—5分"],
    ["反馈帮助率","有帮助反馈数÷有效反馈数","反映用户是否认为结果有帮助"],
    ["再次使用意愿率","愿意再次使用反馈数÷有效反馈数","作为复用倾向的代理指标"],
    ["有效反馈数","有效反馈字段为1的记录数","排除模拟的无效或重复反馈"],
], [2200,3900,3260])
add_heading(doc, "6.3 指标解释边界", 2)
add_body(doc, "生成成功率仅表示模拟事件是否返回结果，不等同于回答准确性；任务完成率反映用户目标是否完成，不等同于模型调用成功；平均评分和帮助率来自反馈子样本，不能直接推断全部模拟用户；异常记录为抽样风险样本，不等同于全部失败事件。")

add_heading(doc, "七、模拟参数与生成规则", 1)
add_heading(doc, "7.1 全局参数", 2)
add_table(doc,["参数","设置值","设置目的"],[
    ["随机种子","20260815","保证相同版本可重复生成"],
    ["模拟周期",SUMMARY["simulatedPeriod"],"覆盖一个连续30天观察窗口"],
    ["用户规模",SUMMARY["users"],"满足分群比较的基本样本量"],
    ["会话规模",SUMMARY["sessions"],"模拟多次访问与连续使用"],
    ["使用事件",SUMMARY["events"],"支持按功能和用户群体统计"],
    ["反馈样本",SUMMARY["feedback"],"支持满意度和再次使用意愿分析"],
    ["异常样本",SUMMARY["anomalies"],"支持异常分级和处置状态复盘"],
], [2300,2200,4860])
add_heading(doc, "7.2 功能使用生成规则", 2)
add_body(doc, "每条事件先从模拟会话中继承用户ID和时间，再结合用户偏好功能与全局功能权重确定功能类型。约34%的事件优先使用用户偏好功能，其余事件按会议纪要、文档问答、任务拆解、周报生成、自然语言咨询、智能推荐和任务自动化的基础权重抽样，以形成相对合理但不完全均衡的功能分布。")
add_heading(doc, "7.3 成功、耗时与完成规则", 2)
add_body(doc, "不同功能设置不同的模拟成功率和响应时间基线，AI接受度对成功概率作小幅修正。成功事件再根据功能完成概率生成任务完成、导出、保存、来源有效性、多轮对话、推荐采纳和状态更新等字段；失败事件根据功能类型生成相应错误类型。")
add_caption(doc, "表3  功能参数范围（模拟设定）")
add_table(doc,["功能","基础成功概率","响应时间基线","主要失败场景"],[
    ["会议纪要","95%","约7秒","字段缺失、接口失败、超时"],
    ["文档问答","92%","约8秒","来源定位不足、解析失败、接口失败"],
    ["任务拆解","90%","约8.5秒","字段缺失、接口失败、超时"],
    ["周报生成","91%","约7.5秒","历史资料不足、导出格式异常、接口失败"],
    ["自然语言咨询","93%","约5.5秒","接口失败、上下文丢失、超时"],
    ["智能推荐","87%","约5秒","理由不足、未命中、接口失败"],
    ["任务自动化","84%","约9.5秒","Function Calling失败、状态更新失败、接口失败"],
], [2000,1700,1900,3760])
add_callout(doc, "参数声明", "以上概率和时间仅为模拟生成参数，用于制造功能差异和验证分析流程，不应表述为真实产品指标、百度官方数据或行业权威平均值。", "FFF2CC")

add_heading(doc, "八、异常与边界场景设计", 1)
add_body(doc, "为了使复盘不仅包含正常流程，还能覆盖第9周风险手册和第10周测试发现的问题，数据集中设置失败事件和100条异常样本。异常记录从失败或高耗时事件中抽样，并记录错误类型、等级、处理状态和处置说明。")
add_table(doc,["场景类别","模拟条件","可能影响","复盘方向"],[
    ["接口异常","接口调用失败或Function Calling失败","生成失败、流程中断","检查降级、重试和错误提示"],
    ["文档解析异常","PDF/DOCX/PPTX/TXT/MD解析失败","无法提取正文或来源","检查格式支持、解析日志和提示"],
    ["来源不足","文档问答来源有效性不足","答案依据不可靠","优化引用定位和“无法确定”策略"],
    ["字段缺失","纪要或任务拆解缺少负责人、时间等字段","结果可执行性下降","统一缺失字段标注规则"],
    ["历史资料不足","周报缺少历史任务或会议输入","内容不完整","提示补充资料并限制确定性表达"],
    ["状态更新失败","任务自动化生成后状态未更新","任务闭环中断","检查状态写回和进度逻辑"],
    ["响应时间过长","事件响应时间超过18秒","等待体验下降","分析功能、时段和错误关联"],
], [1900,2600,2400,2460])

add_heading(doc, "九、数据生成流程", 1)
for x in [
    "确定项目功能范围、用户群体和指标口径。",
    "按固定比例生成300名匿名模拟用户，并配置岗位、AI接受度、频率与偏好功能。",
    "在30天窗口内生成1500次模拟会话，配置日期、时间和关联用户。",
    "在会话基础上生成3000条使用事件，按功能基线计算状态、耗时和转化字段。",
    "从成功事件中抽取500条反馈，根据完成状态和响应时间生成评分及反馈标签。",
    "从失败或高耗时事件中抽取100条异常记录，生成等级、状态和处置说明。",
    "执行主键、外键、值域、日期、隐私与汇总公式检查。",
    "导出Excel，并在汇总看板中通过公式计算核心指标和分功能结果。",
]:
    add_number(doc,x)

add_heading(doc, "十、数据质量校验", 1)
add_table(doc,["校验项","校验方法","通过标准","本次结果"],[
    ["主键唯一性","检查用户ID、事件ID、反馈ID、异常ID重复","重复数为0","通过"],
    ["外键完整性","检查事件用户ID、反馈事件ID和异常事件ID是否存在","无法关联数为0","通过"],
    ["数量一致性","核对各表行数与说明文件","与300/3000/500/100一致","通过"],
    ["值域合理性","检查评分、0/1字段、状态枚举和响应时间","均在规定范围","通过"],
    ["日期范围","检查事件和反馈日期","位于模拟周期内","通过"],
    ["逻辑一致性","失败事件不得标记任务完成；不适用字段不强行赋值","无明显逻辑冲突","通过"],
    ["隐私检查","检查姓名、手机号、身份证号、地址和文档正文","不包含真实敏感信息","通过"],
    ["公式错误扫描","扫描#REF!、#DIV/0!、#VALUE!等错误","错误数为0","通过"],
], [1750,3200,2850,1560])
add_body(doc, "质量校验只能证明数据结构和生成逻辑符合本次设定，不能证明模拟参数与真实用户完全一致。进入真实试点后，应使用相同字段和口径接入经授权、脱敏和合规处理的真实数据，再进行外部有效性验证。")

add_heading(doc, "十一、数据集结果概览", 1)
add_callout(doc, "结果概览", f"本次模拟数据共包含{SUMMARY['events']}条使用事件，模拟生成成功率为{SUMMARY['successRate']:.1f}%，平均响应时间为{SUMMARY['avgResponseSeconds']:.1f}秒，任务完成率为{SUMMARY['completionRate']:.1f}%，有效反馈平均评分为{SUMMARY['avgScore']:.2f}分。以上结果仅用于验证指标口径和复盘方法。", LIGHT_BLUE)
doc.add_picture(str(CHART), width=Inches(6.35))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
add_caption(doc, "图1  模拟数据汇总看板（指标由Excel公式计算）")
rows = []
for fn in SUMMARY["functionCounts"]:
    rows.append([fn, SUMMARY["functionCounts"][fn], f"{SUMMARY['functionSuccessRates'][fn]:.1f}%"])
add_caption(doc, "表4  分功能模拟结果")
add_table(doc,["功能","使用事件数","生成成功率"],rows,[3600,2600,3160])
add_heading(doc, "11.1 初步观察", 2)
for x in [
    "会议纪要和文档问答的模拟使用量较高，可作为第11周复盘的重点核心场景。",
    "任务自动化和智能推荐的模拟成功率相对较低，便于验证异常分析、降级策略和迭代优先级。",
    "功能成功率、任务完成率和用户反馈应分开解读，避免仅凭“返回成功”判断结果质量。",
    "后续复盘应继续按用户群体、使用频率、输入方式和错误类型拆分，避免总体均值掩盖局部问题。",
]:
    add_bullet(doc,x)

add_heading(doc, "十二、复盘使用方法", 1)
add_heading(doc, "12.1 推荐分析顺序", 2)
for x in [
    "总体检查：确认用户数、事件数、成功率、耗时、评分和异常数量。",
    "用户分群：比较四类用户的使用频率、偏好功能、成功率、完成率和反馈。",
    "功能分析：对七项功能分别查看使用量、成功率、耗时、完成率和错误类型。",
    "场景分析：按上传文件、粘贴文本、直接输入以及资料类型拆分。",
    "转化分析：查看生成成功后是否完成、导出、保存、采纳或再次使用。",
    "异常分析：按功能、错误类型、等级和处理状态识别集中问题。",
    "形成建议：将问题归入产品功能、AI融合、用户体验和产品增长四个优化模块。",
]: add_number(doc,x)
add_heading(doc, "12.2 分析输出模板", 2)
add_table(doc,["分析模块","需要回答的问题","建议输出"],[
    ["用户群体","哪类用户使用多、成功率低或评分低？","群体特征、问题和优化动作"],
    ["使用场景","哪种输入方式、资料类型或功能组合表现较弱？","场景差异和流程优化"],
    ["AI能力","失败、来源不足、字段缺失集中在哪里？","Prompt、RAG、降级和安全建议"],
    ["用户体验","耗时、导出、保存、反馈环节是否顺畅？","交互、提示和路径优化"],
    ["产品增长","哪些功能具备较高复用和推荐潜力？","激活、留存和功能引导建议"],
], [1800,4300,3260])

add_heading(doc, "十三、局限性与风险控制", 1)
add_table(doc,["局限/风险","具体表现","控制方式"],[
    ["参数主观性","基础概率和权重来自项目假设而非真实用户统计","标注模拟参数，后续用真实试点数据校准"],
    ["样本偏差","四类用户和七项功能的比例是人为设置","进行多组参数敏感性测试"],
    ["相关性简化","真实行为受到组织、任务难度、文件质量等更多因素影响","后续增加场景难度和上下文变量"],
    ["反馈偏差","反馈只来自部分成功事件，可能高估满意度","在下一版本加入失败反馈和未反馈用户"],
    ["异常抽样","100条异常记录不是全部失败事件","异常率计算应回到使用事件表"],
    ["外推风险","模拟结果不能代表真实市场和真实业务","禁止用于外部披露或商业承诺"],
], [1900,3650,3810])
add_heading(doc, "13.1 版本管理", 2)
add_body(doc, "当用户比例、功能权重、成功概率、响应时间、错误类型或指标口径发生变化时，应更新版本号、随机种子和参数表，重新生成全部记录并保留变更说明。不同版本之间比较时，优先保持用户规模、模拟周期和随机种子不变，以便判断优化参数对结果的影响。")

add_heading(doc, "十四、结论", 1)
add_body(doc, f"本报告完成了百度办公 AI 助手 MVP 第11周模拟用户数据集的完整构建说明。配套Excel形成了由{SUMMARY['users']}名模拟用户、{SUMMARY['sessions']}次会话、{SUMMARY['events']}条使用事件、{SUMMARY['feedback']}条反馈和{SUMMARY['anomalies']}条异常记录组成的关联数据体系，覆盖七项产品功能和四类用户群体。")
add_body(doc, "数据集能够支持总体、用户群体、功能、场景、转化和异常六类复盘，并通过结构化字段与公式汇总验证第9周指标体系和第10周数据看板的主要口径。其价值在于验证分析方法和优化流程，而不是替代真实用户调研或真实运营数据。后续应基于本数据集完成模拟数据复盘，并将结果转化为产品功能、AI融合、用户体验和产品增长四个方向的迭代方案。")

doc.add_page_break()
add_heading(doc, "附录A：字段清单", 1)
add_table(doc,["数据表","字段"],[
    ["模拟用户","用户ID、用户群体、岗位层级、AI接受度、使用频率、偏好功能、模拟注册日期、活跃天数、偏好输入方式、主要设备"],
    ["使用事件","事件ID、用户ID、会话ID、事件时间、功能类型、输入方式、资料类型、生成状态、响应时间、任务完成、导出、保存、来源有效性、多轮对话、推荐采纳、推荐修改、Function Calling、状态更新、错误类型"],
    ["用户反馈","反馈ID、事件ID、用户ID、反馈时间、评分、是否有帮助、再次使用意愿、反馈标签、反馈摘要、有效反馈"],
    ["异常记录","异常ID、事件ID、发生时间、功能类型、错误类型、等级、处理状态、处置说明"],
], [1900,7460])

add_heading(doc, "附录B：项目参考材料", 1)
for x in [
    "第3周：《目标垂类用户全维度画像报告》《用户需求痛点分析与优先级排序报告》《项目核心产品方向与指标体系白皮书》",
    "第7周：《AI产品需求文档》《Figma高保真交互原型设计说明》《自我需求评审记录与优化报告》",
    "第9周：《项目全流程精细化数据指标体系手册》《产品落地风险管控与应急预案手册》",
    "第10周：《项目核心数据看板原型设计》《数据复盘机制规范》《MVP原型功能测试报告》《AI能力效果专项测试报告》《竞品对比测试分析报告》",
]: add_bullet(doc,x)
add_body(doc, "说明：上述文件为本项目内部参考材料，用于保持用户、功能、指标和风险口径一致；模拟数据中的数值不来源于真实百度用户数据库。", after=0)

# Core properties
doc.core_properties.title = "模拟用户数据集构建说明"
doc.core_properties.subject = "百度办公AI助手MVP第11周交付物"
doc.core_properties.author = "项目组"
doc.core_properties.keywords = "模拟用户数据, 数据复盘, 百度办公AI助手, MVP"

doc.save(DOCX)
print(json.dumps({"docx": str(DOCX), "embedded_summary": str(CHART), "pages_expected": "约16-20页"}, ensure_ascii=False, indent=2))
