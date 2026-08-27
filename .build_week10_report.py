from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT = r"D:\PycharmProjects\PythonProject6\第十周项目周报.docx"
NAVY = "17365D"
BLUE = "2E74B5"
LIGHT = "EAF1F8"
GRAY = "666666"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_col_width(cell, dxa):
    cell.width = Inches(dxa / 1440)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_col_width(cell, width)


def format_run(run, size=11, bold=False, color="000000"):
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        format_run(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        format_run(r2)
    else:
        r = p.add_run(text)
        format_run(r)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    for r in p.runs:
        format_run(r)
    if not p.runs:
        format_run(p.add_run(text))
    else:
        p.runs[0].text = text
    return p


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(0.78)
sec.bottom_margin = Inches(0.72)
sec.left_margin = Inches(0.82)
sec.right_margin = Inches(0.82)
sec.header_distance = Inches(0.35)
sec.footer_distance = Inches(0.35)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15

for name, size, color, before, after in [
    ("Heading 1", 16, NAVY, 14, 7),
    ("Heading 2", 13, BLUE, 10, 5),
    ("Heading 3", 11.5, NAVY, 8, 4),
]:
    s = styles[name]
    s.font.name = "Microsoft YaHei"
    s._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = RGBColor.from_string(color)
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)

header = sec.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = header.add_run("百度办公AI助手 MVP｜第10周")
format_run(r, size=9, color="7A8491")

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run("第10周项目周报")
format_run(r, size=9, color="8A8F98")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(12)
p.paragraph_format.space_after = Pt(5)
r = p.add_run("第十周项目周报")
format_run(r, size=24, bold=True, color=NAVY)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(16)
r = p.add_run("数据看板设计与多维度自我测试")
format_run(r, size=13, color=BLUE)

meta = doc.add_table(rows=2, cols=3)
meta.alignment = WD_TABLE_ALIGNMENT.CENTER
meta.style = "Table Grid"
set_table_geometry(meta, [3120, 3120, 3120])
meta_vals = [
    ("项目名称", "百度办公AI助手MVP"),
    ("项目周次", "第10周"),
    ("报告日期", "2026年8月13日"),
]
for col, (label, value) in enumerate(meta_vals):
    for row in meta.rows:
        set_cell_margins(row.cells[col])
        row.cells[col].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(meta.cell(0, col), LIGHT)
    p0 = meta.cell(0, col).paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_run(p0.add_run(label), size=9.5, bold=True, color=NAVY)
    p1 = meta.cell(1, col).paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_run(p1.add_run(value), size=10.5)

doc.add_heading("一、本周工作概述", level=1)
add_body(doc, "本周围绕“数据看板设计与多维度自我测试”开展工作，完成项目核心数据看板原型、标准化数据复盘机制、MVP原型功能验收、AI能力效果专项测试和竞品对比测试。工作重点由单纯验证页面能否运行，进一步扩展到功能完整性、AI输出质量、数据监测口径、异常闭环和市场竞争力分析，形成了较完整的MVP阶段评估材料。")

doc.add_heading("二、本周已完成事项", level=1)
doc.add_heading("2.1 数据看板原型与复盘机制", level=2)
add_bullet(doc, "完成项目核心数据看板原型设计，形成顶部核心指标、核心指标趋势、功能使用分析、用户反馈、异常与风险五个展示区域。")
add_bullet(doc, "明确访问会话数、生成请求数、生成成功率、平均生成时间、功能完成率和用户反馈等核心指标的展示方式。")
add_bullet(doc, "建立每日检查、每周复盘、阶段总结三级复盘机制，明确异常预警、责任分工、问题分级、复测和关闭条件。")

doc.add_heading("2.2 MVP原型功能测试", level=2)
add_bullet(doc, "依据第7周PRD、第9周风险管控手册和数据指标手册，对44项正式验收条目进行自动化、手动操作、源码核验、接口联调和异常场景测试。")
add_bullet(doc, "完成文件上传、文档解析、会议纪要、任务拆解、编辑确认、导出保存、历史记录、用户反馈及基础AI生成流程验证。")
add_bullet(doc, "形成需求追踪表、Bug清单、修复建议和后续优化优先级。")

doc.add_heading("2.3 AI能力效果专项测试", level=2)
add_bullet(doc, "设计并执行50个标准化测试用例，准确性、相关性、安全性、流畅性和有用性五个维度各10个用例。")
add_bullet(doc, "完成安全性和有用性回归测试，验证敏感信息脱敏、高风险请求拒绝以及结构化方案输出能力。")
add_bullet(doc, "形成五维评分结果、主要质量问题清单和后续回归测试方案。")

doc.add_heading("2.4 竞品对比测试", level=2)
add_bullet(doc, "选取WPS AI、Microsoft 365 Copilot和Notion AI作为主要竞品，围绕会议纪要、文档问答、任务拆解和周报生成4个统一场景开展对比。")
add_bullet(doc, "从内容质量、场景适配、操作效率、可编辑性、安全可控性和协作能力等方面分析产品差异。")
add_bullet(doc, "明确本项目轻量化入口、中文办公模板和结构化任务字段方面的优势，以及连续追问、历史数据利用和来源定位方面的不足。")

doc.add_heading("三、本周关键成果", level=1)
table = doc.add_table(rows=1, cols=4)
table.style = "Table Grid"
table.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_geometry(table, [1800, 2520, 1800, 3240])
headers = ["工作模块", "执行规模", "结果", "阶段结论"]
for i, h in enumerate(headers):
    c = table.rows[0].cells[i]
    set_cell_shading(c, NAVY)
    set_cell_margins(c)
    c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_run(p.add_run(h), size=10, bold=True, color=WHITE)
set_repeat_table_header(table.rows[0])
rows = [
    ("功能验收", "44项正式验收", "36项通过、6项不通过、2项部分通过；严格通过率81.8%", "有条件通过，可用于MVP演示和阶段性交付"),
    ("AI能力测试", "50个标准化用例", "45项通过、5项不通过；通过率90%，综合平均分4.16", "达到测试验收标准，总体通过"),
    ("数据看板", "5个页面区域", "完成指标卡、趋势、功能分析、反馈和异常风险设计", "形成详细原型；图中数值为演示数据"),
    ("数据复盘", "三级复盘机制", "覆盖每日检查、每周复盘和阶段总结", "形成问题发现、修复、复测和关闭闭环"),
    ("竞品分析", "3个竞品、4个场景", "完成统一场景对比和差异化定位分析", "任务拆解优势突出，周报历史利用仍需优化"),
]
for ridx, vals in enumerate(rows, start=1):
    cells = table.add_row().cells
    for i, val in enumerate(vals):
        set_cell_margins(cells[i])
        cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in (0, 1) else WD_ALIGN_PARAGRAPH.LEFT
        format_run(p.add_run(val), size=9.5)
        if ridx % 2 == 0:
            set_cell_shading(cells[i], "F7F9FC")
set_table_geometry(table, [1800, 2520, 1800, 3240])

doc.add_heading("四、存在问题与风险", level=1)
add_bullet(doc, "文档问答结果页尚未形成基于当前资料的连续追问闭环，来源定位粒度仍可进一步细化。")
add_bullet(doc, "周报生成尚未稳定读取历史任务和会议纪要，历史资料不足时缺少专门的补充提示。")
add_bullet(doc, "智能推荐部分行为埋点及专项指标展示仍不完整，影响推荐采用率等指标的准确统计。")
add_bullet(doc, "AI输出在统计计算、复杂标准比较、原因分析和Bug判断等场景中仍存在只复述资料、分析不足的问题。")
add_bullet(doc, "竞品侧主要依据公开资料和既有研究进行统一场景适配度走查，未在相同付费账户和模型版本下完成全量盲测，结论主要用于定位与差距识别。")

doc.add_heading("五、下周工作计划", level=1)
plans = [
    ("P0", "完善核心内容闭环", "补充文档问答连续追问、周报历史数据读取和历史不足提示；加强来源精确定位。", "核心场景能够连续完成，缺失资料有明确提示"),
    ("P0", "修复遗留验收问题", "根据44项验收结果修复未通过和部分通过条目，并使用原场景开展回归测试。", "问题状态、修复记录和复测结论可追踪"),
    ("P1", "完善数据采集与指标", "补充推荐行为、保存率和专项功能指标埋点，统一看板与复盘规范中的统计口径。", "关键指标能够由真实事件计算，不使用演示数据替代"),
    ("P1", "提升AI复杂任务能力", "优化统计计算、深度分析、多标准比较和Bug判断类提示词与校验逻辑。", "固定同一版本重新执行相关失败用例"),
    ("P2", "整理阶段交付材料", "统一5份第10周交付物中的术语、结论、数据和章节格式，形成可提交版本。", "报告之间口径一致、引用清楚、无相互矛盾"),
]
pt = doc.add_table(rows=1, cols=4)
pt.style = "Table Grid"
pt.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_geometry(pt, [900, 2100, 3900, 2460])
for i, h in enumerate(["优先级", "工作任务", "主要内容", "验收标准"]):
    c = pt.rows[0].cells[i]
    set_cell_shading(c, NAVY)
    set_cell_margins(c)
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_run(p.add_run(h), size=10, bold=True, color=WHITE)
set_repeat_table_header(pt.rows[0])
for vals in plans:
    cells = pt.add_row().cells
    for i, val in enumerate(vals):
        set_cell_margins(cells[i])
        cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i < 2 else WD_ALIGN_PARAGRAPH.LEFT
        format_run(p.add_run(val), size=9.5)
set_table_geometry(pt, [900, 2100, 3900, 2460])

doc.add_heading("六、需要协调的事项", level=1)
add_bullet(doc, "需要保持稳定的文心千帆接口测试环境，用于功能回归和响应时间复测。")
add_bullet(doc, "后续试运营阶段需要积累一定数量的真实匿名事件和反馈样本，验证看板指标及预警阈值是否合理。")
add_bullet(doc, "AI能力正式复测建议增加第二名评审者进行盲评，以降低单人评分偏差。")

doc.add_heading("七、本周交付物", level=1)
deliverables = [
    "《项目核心数据看板原型设计》",
    "《数据复盘机制规范》",
    "《MVP原型功能测试报告》",
    "《AI能力效果专项测试报告》",
    "《竞品对比测试分析报告》",
]
for item in deliverables:
    add_bullet(doc, item)

doc.add_heading("八、本周总结", level=1)
add_body(doc, "本周完成了百度办公AI助手MVP从功能验收、AI效果评估、数据监测到竞品定位的系统性检查。当前系统核心办公内容生成闭环已基本可用，AI能力专项测试达到阶段验收要求，数据看板与复盘机制能够为后续问题跟踪和版本迭代提供基础。现阶段仍需围绕文档问答连续追问、历史数据驱动周报、专项指标采集和复杂任务分析能力继续优化。下一阶段将以遗留问题回归、指标真实接入和核心场景闭环为重点，进一步提升产品稳定性、可用性和数据可评估性。")

doc.core_properties.title = "第十周项目周报"
doc.core_properties.subject = "数据看板设计与多维度自我测试"
doc.core_properties.keywords = "百度办公AI助手,MVP,第十周,项目周报"
doc.save(OUT)
print(OUT)
