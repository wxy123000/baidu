from docx import Document
p=r'D:\PycharmProjects\PythonProject6\.doc_review_data_retro.docx'
doc=Document(p)
print('PARAGRAPHS',len(doc.paragraphs),'TABLES',len(doc.tables),'SECTIONS',len(doc.sections),'IMAGES',len(doc.inline_shapes))
print('\n=== PARAGRAPHS ===')
for i,x in enumerate(doc.paragraphs):
    t=' '.join(x.text.split())
    if t: print(f'P{i:03d}\t[{x.style.name}]\t{t}')
print('\n=== TABLES ===')
for ti,t in enumerate(doc.tables):
    print(f'\nTABLE {ti+1}: {len(t.rows)}x{len(t.columns)}')
    for ri,row in enumerate(t.rows):
        vals=[' '.join(c.text.split()) for c in row.cells]
        print(f'R{ri:02d}\t'+' | '.join(vals))
