from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\PycharmProjects\PythonProject6")
OUT = ROOT / "outputs" / "week11_simulated_dataset"
DATA = json.loads((OUT / "review_analysis.json").read_text(encoding="utf-8"))
DOCX = OUT / "项目AI产品全案优化迭代方案.docx"

NAVY = "1F4E78"
BLUE = "4472C4"
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "667085"
INK = "1F2937"
GREEN = "70AD47"
AMBER = "ED7D31"
RED = "C00000"
WHITE = "FFFFFF"


def shade(cell, color):
    pr = cell._tc.get_or_add_tcPr()
    node = pr.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        pr.append(node)
    node.set(qn("w:fill"), color)


def cell_width(cell, dxa):
    pr = cell._tc.get_or_add_tcPr()
    node = pr.find(qn("w:tcW"))
    if node is None:
        node = OxmlElement("w:tcW")
        pr.append(node)
    node.set(qn("w:w"), str(dxa))
    node.set(qn("w:type"), "dxa")


def cell_margins(cell, top=80, start=120, bottom=80, end=120):
    pr = cell._tc.get_or_add_tcPr()
    mar = pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        pr.append(mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def fix_table(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    pr = table._tbl.tblPr
    layout = pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tw = pr.find(qn("w:tblW"))
    if tw is None:
        tw = OxmlElement("w:tblW")
        pr.append(tw)
    tw.set(qn("w:w"), str(sum(widths)))
    tw.set(qn("w:type"), "dxa")
    ind = pr.find(qn("w:tblInd"))
    if ind is None:
        ind = OxmlElement("w:tblInd")
        pr.append(ind)
    ind.set(qn("w:w"), "120")
    ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for value in widths:
        node = OxmlElement("w:gridCol")
        node.set(qn("w:w"), str(value))
        grid.append(node)
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
        for i, cell in enumerate(row.cells):
            cell_width(cell, widths[i])
            cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_font(run, size=11, bold=False, color=INK, italic=False):
    run.font.name = "Microsoft YaHei"
    fonts = run._element.get_or_add_rPr().rFonts
    for part in ("eastAsia", "ascii", "hAnsi"):
        fonts.set(qn(f"w:{part}"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_para(p, before=0, after=6, line=1.10, align=None):
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if align is not None:
        p.alignment = align


def add_body(doc, text, after=6):
    p = doc.add_paragraph()
    set_para(p, after=after)
    set_font(p.add_run(text))
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    set_para(p, after=8, line=1.167)
    set_font(p.add_run(text))
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    set_para(p, after=8, line=1.167)
    set_font(p.add_run(text))
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    sizes = {1: 16, 2: 13, 3: 12}
    set_font(p.add_run(text), size=sizes[level], bold=True, color=NAVY if level < 3 else "1F4D78")
    return p


def add_table(doc, headers, rows, widths, font_size=8.8):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    fix_table(table, widths)
    for i, value in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade(cell, LIGHT_GRAY)
        p = cell.paragraphs[0]
        set_para(p, after=0, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_font(p.add_run(str(value)), size=font_size, bold=True, color=NAVY)
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        tr_pr.append(OxmlElement("w:tblHeader"))
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            p = cells[i].paragraphs[0]
            set_para(p, after=0, line=1.0,
                     align=WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT)
            set_font(p.add_run(str(value)), size=font_size)
        if row_index % 2:
            for cell in cells:
                shade(cell, "FAFBFC")
    fix_table(table, widths)
    return table


def add_callout(doc, title, text, fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    fix_table(table, [9360])
    cell = table.cell(0, 0)
    shade(cell, fill)
    p = cell.paragraphs[0]
    set_para(p, after=2)
    set_font(p.add_run(title + "："), bold=True, color=NAVY)
    set_font(p.add_run(text))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_page_number(p):
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(p.add_run("第 "), size=9, color=MID_GRAY)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    p._p.append(field)
    set_font(p.add_run(" 页"), size=9, color=MID_GRAY)


O = DATA["overall"]
G = DATA["groupMetrics"]
F = DATA["functionMetrics"]

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
styles["Normal"].font.name = "Microsoft YaHei"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
styles["Normal"].font.size = Pt(11)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.10
for level, before, after in ((1, 16, 8), (2, 12, 6), (3, 8, 4)):
    style = styles[f"Heading {level}"]
    style.font.name = "Microsoft YaHei"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)

header = section.header.paragraphs[0]
set_font(header.add_run("百度办公 AI 助手｜第11周优化迭代方案"), size=9, color=MID_GRAY)
add_page_number(section.footer.paragraphs[0])

# Memo masthead
p = doc.add_paragraph()
set_para(p, before=18, after=4)
set_font(p.add_run("PRODUCT ITERATION PLAN"), size=10, bold=True, color=BLUE)
p = doc.add_paragraph()
set_para(p, after=5)
set_font(p.add_run("项目 AI 产品全案优化迭代方案"), size=25, bold=True, color=NAVY)
p = doc.add_paragraph()
set_para(p, after=18)
set_font(p.add_run("基于MVP自我测试与模拟数据复盘的四模块优化方案"), size=13, color=MID_GRAY)
meta = [
    ("项目", "百度办公 AI 助手 MVP"),
    ("方案范围", "产品功能、AI融合、用户体验、产品增长"),
    ("依据", "第10周自我测试 + 第11周模拟数据复盘"),
    ("方案周期", "建议4周完成核心整改并持续验证"),
    ("版本", "V1.0｜2026年8月19日"),
]
for label, value in meta:
    p = doc.add_paragraph()
    set_para(p, after=2)
    set_font(p.add_run(label + "："), bold=True)
    set_font(p.add_run(value))
p = doc.add_paragraph()
set_para(p, before=12, after=14)
pPr = p._p.get_or_add_pPr()
pBdr = OxmlElement("w:pBdr")
bottom = OxmlElement("w:bottom")
bottom.set(qn("w:val"), "single")
bottom.set(qn("w:sz"), "12")
bottom.set(qn("w:space"), "1")
bottom.set(qn("w:color"), BLUE)
pBdr.append(bottom)
pPr.append(pBdr)
add_callout(doc, "方案定位", "本方案不是重新设计全部产品，而是以测试和模拟复盘中已经识别的问题为依据，先修复影响核心链路的问题，再优化AI质量、用户体验和增长机制。所有数据结果均为模拟分析结果，不代表百度真实运营表现。", "FFF2CC")
doc.add_page_break()

add_heading(doc, "摘要", 1)
add_body(doc, f"本方案面向百度办公 AI 助手 MVP 的下一轮优化迭代。第10周完成的功能测试、AI能力测试和竞品分析表明，系统已经具备资料输入、会议纪要、文档问答、任务拆解、周报生成、编辑保存、导出、历史记录、用户反馈和数据指标等基础能力；第11周模拟数据复盘进一步显示，总体生成成功率为{O['successRate']:.1f}%，平均响应时间为{O['avgResponseSeconds']:.1f}秒，任务完成率为{O['completionRate']:.1f}%，但任务自动化成功率仅为{F['任务自动化']['successRate']:.1f}%，管理人员群体再次使用意愿仅为{G['管理人员']['reuseIntentRate']:.1f}%。")
add_body(doc, "因此，本轮优化采用“先稳定、再可信、后增长”的顺序：第一阶段解决接口失败、任务状态和关键字段缺失；第二阶段加强文档来源定位、AI安全与可解释性；第三阶段优化完整操作链路和差异化人群体验；第四阶段再通过核心场景运营、反馈闭环和指标看板推动持续使用。")
add_callout(doc, "核心决策", "P0阶段只处理影响任务是否能完成、结果是否可信的问题；非核心视觉美化和复杂增长玩法暂不优先，避免在基础稳定性不足时扩大流量。")

add_heading(doc, "一、方案背景与目标", 1)
add_heading(doc, "1.1 项目现状", 2)
add_body(doc, "当前MVP已形成“资料输入—功能选择—AI生成—人工编辑确认—保存或导出—反馈与数据记录”的基础闭环。系统能够完成核心办公内容生成，但外部接口稳定性、复杂任务结构化质量、来源定位、状态反馈和高价值用户适配仍存在不足。")
add_heading(doc, "1.2 优化目标", 2)
for text in [
    "提升核心功能的稳定性，降低接口失败、生成超时和状态更新失败；",
    "提升AI结果的准确性、字段完整性、来源可追溯性和安全性；",
    "完善从输入到保存、导出和反馈的端到端体验；",
    "围绕核心用户和高价值场景建立可复用、可衡量的增长闭环；",
    "将每一项改动落实为可测试、可验收、可复盘的版本任务。",
]: add_bullet(doc, text)
add_heading(doc, "1.3 方案原则", 2)
add_table(doc, ["原则", "具体要求"], [
    ["问题驱动", "每项优化都必须对应测试问题、模拟指标或用户反馈，不做无依据堆功能"],
    ["核心链路优先", "先保证输入、生成、编辑、保存和导出，再扩展推荐和自动化能力"],
    ["人工可控", "AI结果允许编辑、确认、拒绝、改选和回退，不直接替用户执行高风险动作"],
    ["安全合规", "敏感信息脱敏，高风险请求拒绝，数据事件可追踪但不保存文档正文"],
    ["数据闭环", "需求、埋点、测试、上线、复盘使用统一指标口径"],
], [1900, 7460])

add_heading(doc, "二、问题基线与优先级", 1)
add_heading(doc, "2.1 当前基线", 2)
add_table(doc, ["观察维度", "当前表现", "主要判断"], [
    ["整体稳定性", f"成功率{O['successRate']:.1f}%，平均响应{O['avgResponseSeconds']:.1f}秒", "具备继续迭代基础，但失败仍影响关键场景"],
    ["任务自动化", f"成功率{F['任务自动化']['successRate']:.1f}%，响应{F['任务自动化']['avgResponseSeconds']:.1f}秒", "稳定性和状态一致性不足"],
    ["任务拆解", f"成功率{F['任务拆解']['successRate']:.1f}%，完成率{F['任务拆解']['completionRate']:.1f}%", "字段缺失影响任务可执行性"],
    ["文档问答", f"成功率{F['文档问答']['successRate']:.1f}%，再用意愿{F['文档问答']['reuseIntentRate']:.1f}%", "价值较高，但解析和来源定位需加强"],
    ["管理人员", f"成功率{G['管理人员']['successRate']:.1f}%，再用意愿{G['管理人员']['reuseIntentRate']:.1f}%", "复杂管理场景适配不足"],
    ["异常处理", "100条异常中62条仍在处理或待处理", "需要明确闭环责任和验收状态"],
], [1800, 3000, 4560])
add_heading(doc, "2.2 优先级定义", 2)
add_table(doc, ["级别", "判断标准", "处理要求"], [
    ["P0", "阻断核心流程、可能造成结果失真或安全风险", "进入最近版本，修复后必须回归测试"],
    ["P1", "不阻断流程但明显影响可信度、完成率或高价值用户", "核心稳定后紧接处理"],
    ["P2", "影响效率、体验一致性或持续使用", "纳入体验与增长迭代"],
    ["P3", "视觉细节、低频增强或探索型能力", "根据资源与数据验证结果安排"],
], [1100, 4250, 4010])
add_heading(doc, "2.3 问题清单", 2)
add_table(doc, ["编号", "级别", "问题", "影响模块", "拟解决版本"], [
    ["OPT-01", "P0", "接口失败与超时缺少完整重试和降级闭环", "AI融合/系统稳定性", "V1.1"],
    ["OPT-02", "P0", "任务自动化生成、执行和状态更新边界不清", "产品功能/AI融合", "V1.1"],
    ["OPT-03", "P0", "敏感信息和高风险操作需要统一安全策略", "AI安全/数据安全", "V1.1"],
    ["OPT-04", "P1", "任务拆解负责人、时间、验收标准等字段可能缺失", "产品功能", "V1.2"],
    ["OPT-05", "P1", "文档问答缺少可靠来源定位和资料不足判断", "AI融合", "V1.2"],
    ["OPT-06", "P1", "管理人员与项目负责人场景模板不足", "用户体验", "V1.2"],
    ["OPT-07", "P2", "加载、失败、降级、保存和导出提示不统一", "用户体验", "V1.3"],
    ["OPT-08", "P2", "推荐理由、采纳与改选数据闭环不足", "增长/智能推荐", "V1.3"],
], [950, 700, 3600, 2210, 1900], font_size=8.3)

add_heading(doc, "三、总体优化框架", 1)
add_body(doc, "本轮迭代将产品分为四个相互衔接的模块。产品功能解决“能不能完成工作”；AI融合解决“结果是否稳定、可信和安全”；用户体验解决“用户是否容易理解和操作”；产品增长解决“用户是否愿意持续使用并形成复用”。")
add_table(doc, ["模块", "核心目标", "关键抓手", "核心指标"], [
    ["产品功能", "保证核心办公任务可完成", "结构化字段、编辑确认、保存导出、状态跟踪", "任务完成率、保存率、导出成功率"],
    ["AI融合", "保证生成稳定、可信、安全", "重试降级、RAG引用、安全分类、可解释推荐", "真实模型成功率、结果返回率、来源有效率"],
    ["用户体验", "降低操作和理解成本", "分步流程、状态反馈、示例引导、差异化模板", "完成时长、放弃率、满意度"],
    ["产品增长", "推动激活、留存和核心功能复用", "核心场景入口、模板复用、反馈闭环、分群运营", "激活率、复用率、再用意愿、推荐采纳率"],
], [1500, 2500, 3260, 2100], font_size=8.4)

add_heading(doc, "四、产品功能优化方案", 1)
add_heading(doc, "4.1 核心流程优化", 2)
for text in [
    "统一资料输入：文件上传和文本粘贴进入同一资料状态，清楚展示文件名、大小、格式、解析状态和失败原因。",
    "统一功能选择：用户从首页卡片、侧边栏或上传后的下一步进入功能页时，选中状态与目标功能保持一致。",
    "统一结果处理：所有生成页提供编辑、确认、重新生成、保存、导出和反馈入口，减少不同功能间操作差异。",
    "统一历史复用：历史记录支持查看、搜索、筛选、删除和再次使用，并明确当前MVP的会话级保存边界。",
]: add_bullet(doc, text)
add_heading(doc, "4.2 任务拆解优化", 2)
add_table(doc, ["优化项", "方案", "验收标准"], [
    ["字段完整性", "固定任务名称、说明、负责人、优先级、截止时间、依赖、验收标准和风险", "必填字段完整率≥95%"],
    ["未知信息处理", "资料未提及时统一标记“待确认”，禁止自行补充具体人名和日期", "无依据确定性信息为0"],
    ["可执行性", "支持编辑字段、调整顺序、勾选完成和导出", "编辑保存、进度更新和导出均通过"],
    ["复杂任务", "将大任务拆成阶段与子任务，显示依赖关系", "至少支持两级任务结构"],
], [1900, 4500, 2960])
add_heading(doc, "4.3 任务自动化优化", 2)
add_body(doc, "任务自动化不应把“生成任务清单”与“在外部平台真实执行任务”混为一谈。本轮仅保证结构化任务生成、状态跟踪和辅助规划；任何外部执行必须经过用户确认、权限校验和操作日志记录。")
for text in [
    "将状态拆分为待确认、待执行、执行中、已完成、失败和已取消；",
    "勾选与取消操作必须保证进度计算一致，避免状态反向累计；",
    "Function Calling失败时保留输入内容，提供重试、本地模拟任务和手动创建三种选择；",
    "高风险操作如删除全部数据、绕过权限或外发机密文件必须拒绝自动执行。",
]: add_bullet(doc, text)
add_heading(doc, "4.4 文档问答与内容生成优化", 2)
add_table(doc, ["功能", "优化方向", "关键结果"], [
    ["文档问答", "增加文档名、页码/段落、引用片段和资料不足提示", "回答有依据且可定位"],
    ["会议纪要", "保持摘要、结论、待办、负责人、时间和风险结构", "未知字段统一待确认"],
    ["周报生成", "区分已完成、问题、下周计划与协调事项，避免把计划写成成果", "信息归类正确"],
    ["智能推荐", "展示推荐理由、适用条件和替代选项", "支持采纳、修改和拒绝"],
], [1700, 4750, 2910])

add_heading(doc, "五、AI融合优化方案", 1)
add_heading(doc, "5.1 模型调用与降级机制", 2)
add_number(doc, "用户发起生成请求后，系统先完成输入安全分类、文件解析和字段校验。")
add_number(doc, "调用千帆模型并记录request_sent、generation_succeeded或generation_failed事件。")
add_number(doc, "对短暂超时进行有限次数重试；连续失败时触发本地模板降级。")
add_number(doc, "本地结果成功展示时记录local_fallback_used和result_displayed，避免看板仍显示0%成功。")
add_number(doc, "如果模型与本地降级均失败，展示明确错误原因、保留用户输入并提供重试。")
add_heading(doc, "5.2 指标口径优化", 2)
add_table(doc, ["指标", "建议定义", "目标"], [
    ["真实模型成功率", "千帆返回可用结果次数÷千帆请求次数", "≥92%"],
    ["本地降级率", "触发并成功展示本地结果次数÷全部生成请求", "持续监控，不以越高越好"],
    ["总体结果返回率", "模型成功或降级成功次数÷全部生成请求", "≥98%"],
    ["平均生成时间", "从发起请求到页面显示最终结果的平均时间", "≤10秒为理想，10—30秒需优化"],
    ["来源有效率", "可定位到有效文档片段的问答数÷文档问答数", "≥90%"],
], [2000, 5000, 2360])
add_heading(doc, "5.3 RAG与来源可信度", 2)
for text in [
    "文档解析阶段保存分段编号、文档名、页码或幻灯片编号等元数据；",
    "回答必须基于检索到的片段，相关度不足时明确说明“资料中未提及”；",
    "结果页显示简短关键依据，不复制过长原文，不显示完整敏感信息；",
    "对跨文档冲突标明来源差异，并提示用户确认采用哪一条标准。",
]: add_bullet(doc, text)
add_heading(doc, "5.4 安全与合规", 2)
add_table(doc, ["风险类别", "系统策略", "示例"], [
    ["敏感信息", "识别并脱敏身份证号、手机号、银行卡号、密码和地址", "仅显示必要的部分号码"],
    ["越权访问", "拒绝获取密码、绕过权限、未经授权登录和导出", "引导走授权或共享流程"],
    ["破坏性操作", "删除、覆盖和批量外发前要求审批、备份和范围确认", "优先软删除和可恢复操作"],
    ["歧视与不当建议", "拒绝基于性别等受保护属性进行排除", "改为统一能力评价标准"],
    ["事件隐瞒", "数据泄露和重大异常必须记录、上报并保留证据", "不允许省略记录以通过验收"],
], [1800, 4600, 2960])

add_heading(doc, "六、用户体验优化方案", 1)
add_heading(doc, "6.1 信息架构与导航", 2)
for text in [
    "侧边栏使用唯一页面状态：历史记录、我的文件、数据指标和帮助反馈分别高亮自身，不再默认高亮首页；",
    "首页功能卡片、侧边栏功能入口与上传后的功能选择页使用统一跳转逻辑；",
    "帮助与反馈固定在侧边栏底部，完成生成后增加轻量评分提醒；",
    "数据看板以核心指标、功能表现、用户反馈和异常风险四个区域组织信息。",
]: add_bullet(doc, text)
add_heading(doc, "6.2 状态与错误反馈", 2)
add_table(doc, ["状态", "页面表现", "用户可执行动作"], [
    ["解析中", "显示文件与解析进度", "取消或等待"],
    ["生成中", "显示加载状态和预计等待", "保留页面并等待"],
    ["模型失败/降级成功", "明确说明已使用本地结果", "接受、编辑或重试模型"],
    ["完全失败", "显示原因且保留输入", "重试、重新上传或使用模板"],
    ["保存/导出成功", "展示文件位置或下载入口", "打开文件或返回历史"],
], [1600, 4200, 3560])
add_heading(doc, "6.3 分群体验", 2)
add_table(doc, ["用户群体", "重点场景", "体验优化"], [
    ["职场新人", "咨询、会议纪要、周报", "示例输入、字段解释、模板引导"],
    ["普通职员", "会议纪要、周报、文档问答", "历史复用、快捷导出、周度提醒"],
    ["项目负责人", "文档问答、任务拆解、推荐", "责任人、依赖、风险和来源依据"],
    ["管理人员", "管理摘要、推荐、任务跟踪", "结论优先、风险总览、跨任务状态"],
], [1800, 3000, 4560])

add_heading(doc, "七、产品增长优化方案", 1)
add_heading(doc, "7.1 增长路径", 2)
add_table(doc, ["阶段", "用户目标", "产品动作", "衡量指标"], [
    ["触达", "知道产品能解决什么问题", "用会议纪要和周报展示明确价值", "入口访问率"],
    ["激活", "完成第一次有效生成", "示例资料、推荐功能、低门槛模板", "首次生成成功率"],
    ["留存", "在固定办公场景重复使用", "历史复用、周度提醒、常用模板", "7日/30日复用率"],
    ["深化", "从单功能扩展到组合使用", "问答转任务、纪要转待办、周报引用历史", "跨功能使用率"],
    ["反馈", "主动提交质量评价", "结果页轻量评分和问题标签", "有效反馈率"],
], [1300, 2300, 3500, 2260])
add_heading(doc, "7.2 核心增长场景", 2)
for text in [
    "以会议纪要作为首次体验入口：结构清晰、成功率高、用户容易判断价值；",
    "以周报生成建立周期性使用：支持复用上周内容和本周历史任务；",
    "以文档问答提高资料使用深度：答案可追溯后更容易形成信任；",
    "以任务拆解连接内容生成与行动：将纪要结论转为可执行任务；",
    "任务自动化仅在稳定性、权限和确认机制达标后逐步开放。",
]: add_bullet(doc, text)
add_heading(doc, "7.3 增长实验", 2)
add_table(doc, ["实验", "对照方案", "实验方案", "成功标准"], [
    ["首次激活", "空白输入页", "提供3个示例输入和推荐功能", "首次成功生成率提升≥8%"],
    ["周报留存", "无提醒", "每周固定时间提醒复用上周内容", "周报7日复用率提升≥10%"],
    ["推荐采纳", "只显示推荐结果", "增加推荐理由与替代选项", "采纳率提升≥8%"],
    ["反馈提交", "只在帮助页反馈", "生成后轻量评分提醒", "有效反馈率提升≥5%"],
], [1700, 2350, 3300, 2010], font_size=8.4)

add_heading(doc, "八、版本路线图", 1)
add_table(doc, ["版本", "周期", "重点范围", "主要输出", "发布条件"], [
    ["V1.1 稳定性版", "第1周", "接口重试降级、任务状态、安全策略", "稳定调用链路与统一事件口径", "P0测试全部通过"],
    ["V1.2 可信度版", "第2周", "任务字段、RAG来源、分群模板", "可信结果与可追溯依据", "字段完整率和来源有效率达标"],
    ["V1.3 体验版", "第3周", "导航、状态提示、历史复用、导出", "端到端体验闭环", "核心流程回归通过"],
    ["V1.4 增长验证版", "第4周", "核心场景运营、反馈和增长实验", "实验方案与数据看板", "埋点完整且可复盘"],
], [1600, 1100, 2650, 2500, 1510], font_size=8.2)
add_heading(doc, "8.1 依赖关系", 2)
add_body(doc, "稳定性和安全策略是全部后续工作的前置条件；RAG来源和结构化字段是可信度版的核心；体验优化依赖统一状态和事件口径；增长实验必须在核心流程稳定、埋点准确后开展。")

add_heading(doc, "九、验收与测试方案", 1)
add_heading(doc, "9.1 核心验收指标", 2)
add_table(doc, ["模块", "验收指标", "目标值", "验证方式"], [
    ["系统稳定性", "总体结果返回率", "≥98%", "接口成功与本地降级联合统计"],
    ["任务自动化", "结构化任务生成成功率", "≥90%", "标准任务用例与断网降级测试"],
    ["任务拆解", "关键字段完整率", "≥95%", "抽检负责人、时间、依赖和验收字段"],
    ["文档问答", "来源有效率", "≥90%", "答案与引用片段人工核验"],
    ["安全性", "高风险请求正确拒绝率", "100%", "敏感、越权、破坏性用例回归"],
    ["用户体验", "核心流程可完成率", "≥95%", "上传至保存/导出端到端测试"],
    ["反馈", "反馈成功写入率", "100%", "评分、帮助度、再用意愿和文本测试"],
], [1600, 3000, 1400, 3360])
add_heading(doc, "9.2 测试层级", 2)
for text in [
    "功能回归：按照PRD验收条目验证上传、生成、编辑、保存、导出、历史、反馈和指标；",
    "AI专项测试：沿用准确性、相关性、安全性、流畅性和有用性五维用例；",
    "异常测试：覆盖断网、接口超时、解析失败、字段缺失、导出异常和重复操作；",
    "数据核验：核对事件名称、分母分子、失败原因和看板展示是否一致；",
    "小流量验证：模拟测试通过后再使用合规的真实小样本校准结论。",
]: add_bullet(doc, text)

add_heading(doc, "十、数据埋点与复盘机制", 1)
add_table(doc, ["流程阶段", "建议事件", "关键字段"], [
    ["资料输入", "upload_started / upload_succeeded / upload_failed", "格式、大小、数量、错误类型"],
    ["模型生成", "request_sent / generation_succeeded / generation_failed", "功能、模型、耗时、错误"],
    ["降级返回", "local_fallback_used / result_displayed", "降级类型、最终是否展示"],
    ["结果处理", "result_edited / result_saved / result_exported", "编辑、保存、导出格式"],
    ["推荐", "recommendation_shown / accepted / modified", "推荐功能、理由、用户选择"],
    ["任务", "task_created / status_updated / completed", "负责人、优先级、截止时间、状态"],
    ["反馈", "feedback_submitted", "评分、帮助度、再用意愿、标签"],
], [1800, 3400, 4160], font_size=8.4)
add_body(doc, "复盘频率保持“每日异常检查—每周指标复盘—版本阶段总结”。每个问题必须记录发现依据、优先级、负责人、目标时间、修复版本、回归结果和最终状态，形成闭环。")

add_heading(doc, "十一、风险与应对", 1)
add_table(doc, ["风险", "可能影响", "应对措施"], [
    ["外部模型接口继续不稳定", "核心生成失败、体验波动", "限次重试、本地降级、熔断和监控告警"],
    ["指标口径前后不一致", "成功率和优化效果失真", "统一事件字典、分母分子和版本说明"],
    ["模拟数据与真实行为偏差", "错误判断优先级", "小流量校准，真实结果优先"],
    ["新增功能扩大安全风险", "敏感信息、越权或误执行", "默认最小权限、确认、审计和拒绝策略"],
    ["版本范围过大", "延期或回归不足", "严格按P0/P1拆分，未达标不进入下一阶段"],
], [2200, 3100, 4060])

add_heading(doc, "十二、预期效果", 1)
add_body(doc, "完成本轮优化后，产品应从“能够演示功能”提升为“核心流程可稳定完成、AI结果可解释、用户操作可理解、问题能够被数据发现”的MVP。会议纪要、周报和文档问答继续承担核心价值；任务拆解提升可执行性；任务自动化在权限、确认和稳定性达标后逐步开放；增长工作以真实复用和任务完成为目标，而不是单纯追求访问次数。")
add_table(doc, ["维度", "优化前主要问题", "优化后目标状态"], [
    ["产品功能", "功能可用但字段和状态不完全一致", "核心功能结构统一、状态可靠、可编辑交付"],
    ["AI融合", "接口失败、来源不足、降级统计不准确", "调用可恢复、结果可追溯、安全策略统一"],
    ["用户体验", "导航与状态反馈存在不一致", "操作路径连续、失败可理解、结果易处理"],
    ["产品增长", "以单次生成和模拟反馈为主", "围绕核心场景形成激活、复用和反馈闭环"],
], [1700, 3700, 3960])

add_heading(doc, "十三、结论", 1)
add_body(doc, "本方案将第10周自我测试和第11周模拟数据复盘的发现转化为可实施的产品迭代任务，并按照产品功能、AI融合、用户体验和产品增长四个模块形成完整路线。下一步应先完成V1.1稳定性版，验证接口降级、任务状态和安全策略，再推进可信度、体验和增长优化。所有版本均需通过功能回归、AI专项测试、异常测试和数据核验后进入下一阶段。")
add_callout(doc, "交付结论", "《项目AI产品全案优化迭代方案》已经明确问题、优先级、实施动作、版本路线、验收指标和风险控制，可作为第11周第3项交付物以及第4项MVP原型优化的实施依据。")

add_heading(doc, "附录A：方案依据", 1)
add_table(doc, ["材料", "在本方案中的用途"], [
    ["MVP原型功能测试报告", "确定功能缺陷、异常场景和回归要求"],
    ["AI能力效果专项测试报告", "确定准确性、相关性、安全性、流畅性和有用性问题"],
    ["竞品对比测试分析报告", "识别功能与体验差异化方向"],
    ["项目核心数据看板原型设计", "确定核心指标与页面展示结构"],
    ["数据复盘机制规范", "确定问题发现、整改和回归闭环"],
    ["模拟用户数据集与复盘报告", "确定人群、功能和场景优先级"],
], [3500, 5860])
add_heading(doc, "附录B：实施原则", 1)
for text in [
    "不把模拟数据写成真实运营数据；",
    "不为了提高通过率而修改测试标准；",
    "不把本地降级结果误记为真实模型成功；",
    "不在未授权的情况下执行外部任务或处理敏感数据；",
    "不在核心稳定性未达标前大规模推广。",
]: add_bullet(doc, text)

props = doc.core_properties
props.title = "项目AI产品全案优化迭代方案"
props.subject = "百度办公 AI 助手第11周第3项交付物"
props.author = "AI产品经理实习项目"
props.keywords = "AI产品, 优化迭代, MVP, 产品功能, AI融合, 用户体验, 产品增长"
doc.save(DOCX)
print(DOCX)
