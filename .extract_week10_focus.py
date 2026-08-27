from docx import Document
from pathlib import Path
files=[r'C:\Users\20159\Desktop\实习\第十周\MVP原型功能测试报告.docx',r'C:\Users\20159\Desktop\实习\第十周\竞品对比测试分析报告.docx']
keys=['结论','通过率','通过','不通过','综合','主要问题','优势','不足','建议','限制','风险','完成','测试结果','竞品']
for f in files:
 d=Document(f);print('\n###',Path(f).name)
 for i,p in enumerate(d.paragraphs):
  t=' '.join(p.text.split())
  if t and (p.style.name.startswith('Heading') or any(k in t for k in keys)):
   print(f'P{i:03d}[{p.style.name}] {t}')
 for ti,t in enumerate(d.tables):
  print(f'--TABLE{ti+1}--')
  for ri,row in enumerate(t.rows):
   s=' | '.join(' '.join(c.text.split()) for c in row.cells)
   if ri==0 or any(k in s for k in keys):print(s)
